"""v5.3 Office suite: builders, readers, untrusted-file checks, LibreOffice conversion, lint, tools, server files API."""

import base64
import io
import shutil
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("openpyxl")
pytest.importorskip("pptx")

from automoto import office, office_build, office_check, policy, sandbox, tools  # noqa: E402
from automoto.office import OfficeError  # noqa: E402

HAS_LO = office.lo_available()
HAS_POPPLER = bool(shutil.which("pdftoppm")) and bool(shutil.which("pdfinfo"))
needs_lo = pytest.mark.skipif(not HAS_LO, reason="LibreOffice not installed")
needs_render = pytest.mark.skipif(not (HAS_LO and HAS_POPPLER), reason="LibreOffice/poppler not installed")


@pytest.fixture
def ws():
    root = sandbox.root()
    root.mkdir(parents=True, exist_ok=True)
    return root


@pytest.fixture
def png(ws):
    from PIL import Image

    p = ws / "fig.png"
    Image.new("RGB", (400, 240), (30, 80, 160)).save(p)
    return p


def docx_spec(img=None, toc=True):
    blocks = [{"type": "title_page", "title": "Energy Report", "subtitle": "A test", "author": "Test Author", "date": "2026"}]
    if toc:
        blocks.append({"type": "toc"})
    blocks += [
        {"type": "heading", "level": 1, "text": "Introduction"},
        {"type": "paragraph", "text": "Solar is **growing** fast (Smith, 2020). See [IEA](https://www.iea.org).", "comment": "check figure"},
        {"type": "heading", "level": 2, "text": "Context"},
        {"type": "bullets", "items": ["one", ["nested", 1]]},
        {"type": "numbered", "items": ["first", "second"]},
        {"type": "table", "caption": "Capacity", "rows": [["Year", "GW"], ["2024", "1,200"], ["2025", "1,450"]], "source": "IEA"},
    ]
    if img:
        blocks.append({"type": "image", "path": str(img), "caption": "Blue box", "alt": "A blue rectangle"})
    blocks += [{"type": "page_break"}, {"type": "heading", "level": 1, "text": "Conclusion"},
               {"type": "paragraph", "text": "Done."},
               {"type": "references", "items": [
                   {"authors": ["Smith, J.", "Jones, K."], "year": 2020, "title": "Solar futures", "publisher": "Green Press", "place": "London"},
                   {"authors": "Adams, R.", "year": 2019, "title": "PV costs", "journal": "Energy Policy", "volume": 12, "issue": 3, "pages": "1-9", "doi": "10.1/x"},
                   {"authors": "IEA", "year": 2024, "title": "Renewables 2024", "url": "https://www.iea.org/r", "accessed": "1 May 2026"}]}]
    return {"properties": {"title": "Energy Report", "author": "Test Author"}, "header": "Energy Report", "blocks": blocks}


# ------------------------------------------------------------------ builders + readers
def test_harvard_formats():
    assert office_build.harvard({"authors": ["Smith, J.", "Jones, K."], "year": 2020, "title": "T", "publisher": "P", "place": "L"}) == \
        "Smith, J. and Jones, K. (2020) *T*. L: P."
    j = office_build.harvard({"authors": "Adams, R.", "year": 2019, "title": "PV", "journal": "EP", "volume": 1, "issue": 2, "pages": "3-4"})
    assert j == "Adams, R. (2019) 'PV', *EP*, 1(2), pp. 3-4."
    w = office_build.harvard({"authors": "IEA", "year": 2024, "title": "R", "url": "https://x.org", "accessed": "1 May 2026"})
    assert "Available at: https://x.org (Accessed: 1 May 2026)." in w


def test_docx_build_and_read(ws, png):
    out = ws / "scratch" / "r.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    msg = office_build.create_docx(docx_spec(png, toc=False), out)
    assert "4 headings" in msg and "1 table" in msg and "1 figure" in msg
    text = office.read(str(out))
    assert "# Introduction" in text and "## Context" in text
    assert "| Year | GW |" in text and "Figure 1: Blue box" in text and "Table 1: Capacity" in text
    # references sorted alphabetically, Harvard style
    refs = text.split("# References", 1)[1]
    assert refs.index("Adams") < refs.index("IEA") < refs.index("Smith")
    assert "[comment" in text.lower() or "check figure" in text
    import docx

    d = docx.Document(str(out))
    assert d.core_properties.title == "Energy Report"
    pics = list(d.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"))
    assert pics and pics[0].get("descr") == "A blue rectangle"
    assert any("PAGE" in (t.text or "") for t in d.sections[0].footer._element.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText"))


def test_docx_rejects_bad_blocks_and_images(ws):
    with pytest.raises(OfficeError):
        office_build.build_docx({"blocks": [{"type": "nope"}]}, ws / "x.docx")
    (ws / "a.svg").write_text("<svg/>")
    with pytest.raises(OfficeError):
        office_build.build_docx({"blocks": [{"type": "image", "path": "a.svg"}]}, ws / "x.docx")
    with pytest.raises(sandbox.SandboxError):
        office_build.build_docx({"blocks": [{"type": "image", "path": "/etc/passwd"}]}, ws / "x.docx")


def test_markdown_to_spec():
    spec = office_build.markdown_to_spec("# T\n\nHello *x*\nworld\n\n- a\n  - b\n1. one\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n> q\n")
    kinds = [b["type"] for b in spec["blocks"]]
    assert kinds == ["heading", "paragraph", "bullets", "numbered", "table", "quote"]
    assert spec["blocks"][1]["text"] == "Hello *x* world"
    assert spec["blocks"][2]["items"] == [["a", 0], ["b", 1]]
    assert spec["blocks"][4]["rows"] == [["A", "B"], ["1", "2"]]


XLSX_SPEC = {"sheets": [{"name": "Model", "rows": [["Month", "Revenue", "Cost", "Profit"], ["Jan", 100, 60, "=B2-C2"],
                                                    ["Feb", 120, 70, "=B3-C3"], ["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)"]],
                         "number_formats": {"B2:D4": "£#,##0"}, "bold": ["A4:D4"],
                         "conditional": [{"type": "data_bar", "range": "D2:D3"}, {"type": "cell_is", "range": "D2:D3", "operator": "lessThan", "values": [0]}],
                         "validation": [{"range": "E2", "values": ["High", "Low"]}],
                         "charts": [{"type": "column", "data": "B1:D3", "categories": "A2:A3", "title": "P&L", "anchor": "G2"}]}]}


def test_xlsx_build_and_read(ws):
    out = ws / "m.xlsx"
    msg = office_build.create_xlsx(XLSX_SPEC, out)
    assert "4 rows, 5 formulas, 1 chart" in msg
    import openpyxl

    wb = openpyxl.load_workbook(out)
    s = wb["Model"]
    assert s["D2"].value == "=B2-C2" and s["B2"].number_format == "£#,##0" and s.freeze_panes == "A2"
    assert len(s._charts) == 1 and s.data_validations.dataValidation
    text = office.read(str(out))
    assert "Model" in text and "=SUM(B2:B3)" in text


def test_xlsx_bad_sheet_name(ws):
    with pytest.raises(OfficeError):
        office_build.create_xlsx({"sheets": [{"name": "a/b", "rows": []}]}, ws / "x.xlsx")


PPTX_SPEC = {"theme": {"primary": "#1F3864", "accent": "#2E75B6"}, "slides": [
    {"layout": "title", "title": "Deck", "subtitle": "Sub", "notes": "hello"},
    {"layout": "bullets", "title": "Points", "bullets": ["one", ["two", 1], "**three**"], "notes": "say this"},
    {"layout": "two_content", "title": "Compare", "left": ["a"], "right": ["b"], "notes": "n"},
    {"layout": "table", "title": "Table", "rows": [["A", "B"], ["1", "2"]], "notes": "n"},
    {"layout": "chart", "title": "Chart", "chart": {"type": "column", "categories": ["Q1", "Q2"], "series": [{"name": "Sales", "values": [1, 2]}]}, "notes": "n"},
    {"layout": "statement", "text": "One big idea", "notes": "n"}]}


def test_pptx_build_and_read(ws, png):
    spec = {**PPTX_SPEC, "slides": PPTX_SPEC["slides"] + [{"layout": "image", "title": "Pic", "image": str(png), "caption": "Blue", "notes": "n"}]}
    out = ws / "d.pptx"
    msg = office_build.create_pptx(spec, out)
    assert "7 slides" in msg and "speaker notes on 7" in msg
    text = office.read(str(out))
    assert "13.33×7.50" in text and "## Slide 2: Points" in text and "Notes: say this" in text and "[chart" in text
    import pptx

    prs = pptx.Presentation(str(out))
    for s in prs.slides:  # no empty placeholders left behind
        for ph in s.placeholders:
            assert not ph.has_text_frame or ph.text_frame.text.strip()


def test_edit_docx_xlsx_pptx(ws):
    d = ws / "e.docx"
    office_build.build_docx({"blocks": [{"type": "paragraph", "text": "Hello **world** today"}]}, d)
    assert "1 change" in office_build.edit(d, replace=[{"old": "Hello world", "new": "Hi there"}])
    office_build.edit(d, append=[{"type": "heading", "text": "Added"}])
    t = office.read(str(d))
    assert "Hi there today" in t and "# Added" in t
    x = ws / "e.xlsx"
    office_build.create_xlsx(XLSX_SPEC, x)
    office_build.edit(x, cells={"Model!B2": 500})
    import openpyxl

    assert openpyxl.load_workbook(x)["Model"]["B2"].value == 500
    p = ws / "e.pptx"
    office_build.create_pptx(PPTX_SPEC, p)
    office_build.edit(p, replace=[{"old": "Points", "new": "Key points"}], notes={"1": "new notes"})
    t = office.read(str(p))
    assert "Key points" in t and "Notes: new notes" in t


# ------------------------------------------------------------------ untrusted input
def _zip(path, parts, compress=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(path, "w", compress) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def test_inspect_rejects_zip_bomb(ws):
    p = ws / "bomb.docx"
    _zip(p, {"[Content_Types].xml": "<Types/>", "word/document.xml": b"<a>" + b"0" * (30 * 1024 * 1024) + b"</a>"})
    with pytest.raises(OfficeError, match="zip bomb"):
        office.inspect_file(p)


def test_inspect_rejects_xxe(ws):
    p = ws / "xxe.docx"
    _zip(p, {"[Content_Types].xml": "<Types/>",
             "word/document.xml": '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'})
    with pytest.raises(OfficeError, match="XXE"):
        office.inspect_file(p)
    with pytest.raises(OfficeError):
        office.read(str(p))


def test_inspect_rejects_bad_zip_and_traversal(ws):
    (ws / "bad.xlsx").write_bytes(b"not a zip")
    with pytest.raises(OfficeError, match="bad zip"):
        office.inspect_file(ws / "bad.xlsx")
    p = ws / "trav.pptx"
    _zip(p, {"../evil.xml": "<x/>"})
    with pytest.raises(OfficeError, match="unsafe part"):
        office.inspect_file(p)


def test_inspect_flags_macros_and_external_links(ws):
    d = ws / "m.docx"
    office_build.build_docx({"blocks": [{"type": "paragraph", "text": "see [x](https://example.com/a)"}]}, d)
    info = office.inspect_file(d)
    assert info["external_links"] == ["https://example.com/a"] and not info["macros"]
    m = ws / "macro.docm"
    buf = io.BytesIO()
    with zipfile.ZipFile(d) as src, zipfile.ZipFile(buf, "w") as dst:
        for e in src.infolist():
            dst.writestr(e, src.read(e))
        dst.writestr("word/vbaProject.bin", b"\x00" * 100)
    m.write_bytes(buf.getvalue())
    assert office.inspect_file(m)["macros"] is True


def test_inspect_size_limit(ws, monkeypatch):
    monkeypatch.setattr(office, "MAX_FILE", 10)
    (ws / "big.docx").write_bytes(b"x" * 100)
    with pytest.raises(OfficeError, match="limit"):
        office.inspect_file(ws / "big.docx")


def test_lo_profile_disables_macros():
    xcu = office._XCU
    assert '"MacroSecurityLevel"' in xcu and "<value>3</value>" in xcu
    assert '"DisableMacrosExecution" oor:op="fuse"><value>true' in xcu
    for app in ("Calc", "Writer", "Impress"):
        assert f"/org.openoffice.Office.{app}/Filter/Import/VBA" in xcu
    assert "Content/Update\"><prop oor:name=\"Link\" oor:op=\"fuse\"><value>1" in xcu  # Calc: never update links
    arg = office._filter_arg("pdf", ".docx", None)
    assert arg.startswith("pdf:writer_pdf_Export:") and '"UseTaggedPDF":{"type":"boolean","value":"true"}' in arg
    assert office._filter_arg("pdf", ".pptx", None).startswith("pdf:impress_pdf_Export:")


# ------------------------------------------------------------------ LibreOffice / poppler
@needs_lo
def test_convert_docx_to_pdf_and_toc_pages(ws, png):
    out = ws / "report.docx"
    msg = office_build.create_docx(docx_spec(png), out)
    assert "table of contents filled" in msg and "(4/4 headings found)" in msg
    t = office.read(str(out))
    assert "Introduction\t3" in t and "Conclusion\t4" in t  # title page, contents, then body
    res = office.convert(str(out), "pdf")
    assert "tagged: yes" in res.lower() or not HAS_POPPLER
    assert (ws / "report.pdf").exists()
    if HAS_POPPLER:
        rep = office_check.check(str(ws / "report.pdf"))
        assert "fonts not embedded" not in rep and "not tagged" not in rep


@needs_lo
def test_recalc_reports_values_and_errors(ws):
    out = ws / "m.xlsx"
    spec = {"sheets": [{"name": "S", "rows": [["a", "b", "c"], [10, 0, "=A2/B2"], [5, 2, "=A3*B3"], ["", "", "=Missing!A1"]]}]}
    office_build.create_xlsx(spec, out)
    r = office.recalc(str(out), save=True)
    assert "recalculated 3 formula(s)" in r and "S!C3 =A3*B3 = 10" in r
    assert "#DIV/0!" in r and "error(s)" in r
    assert (ws / "m.recalc.xlsx").exists()
    rep = office_check.check(str(out))
    assert "missing sheet 'Missing'" in rep


@needs_lo
def test_read_legacy_and_odf_via_libreoffice(ws):
    src = ws / "a.docx"
    office_build.build_docx({"blocks": [{"type": "heading", "text": "Legacy test"}]}, src)
    office.convert(str(src), "odt")
    text = office.read(str(ws / "a.odt"))
    assert "Legacy test" in text and "converted from .odt" in text


@needs_lo
def test_markdown_to_pdf(ws):
    (ws / "n.md").write_text("# Notes\n\nSome text.\n\n- a\n- b\n", encoding="utf-8")
    assert "wrote n.docx" in office.convert("n.md", "docx")
    assert "wrote n.pdf" in office.convert("n.md", "pdf")


@needs_render
def test_render_and_pptx_check(ws):
    out = ws / "d.pptx"
    office_build.create_pptx(PPTX_SPEC, out)
    pages = office.render(str(out), 1, 2, 40)
    assert len(pages) == 2 and all(p.suffix == ".png" and p.stat().st_size > 500 for p in pages)
    rep = office_check.check(str(out))
    assert "0 error(s)" in rep and "6 slide(s)" in rep


def test_pptx_check_flags_overflow_and_empty(ws):
    out = ws / "bad.pptx"
    office_build.create_pptx({"slides": [{"layout": "bullets", "title": "Too much", "bullets": ["word " * 60] * 8}]}, out)
    import pptx

    prs = pptx.Presentation(str(out))
    s = prs.slides.add_slide(prs.slide_layouts[1])  # untouched layout → empty placeholders
    prs.save(str(out))
    rep = office_check.check(str(out))
    assert "probably overflows" in rep and "empty placeholder" in rep and "no speaker notes" in rep
    assert s is not None


def test_docx_check_flags_heading_jump_and_alt(ws, png):
    out = ws / "c.docx"
    office_build.build_docx({"blocks": [{"type": "heading", "level": 2, "text": "Start"}, {"type": "heading", "level": 4, "text": "Deep"},
                                        {"type": "image", "path": str(png)}]}, out)
    import docx

    d = docx.Document(str(out))
    for dp in d.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
        dp.set("descr", "")
    d.save(str(out))
    rep = office_check.check(str(out))
    assert "first heading is level 2" in rep and "jumps from 2 to 4" in rep and "no alt text" in rep and "no document title" in rep


# ------------------------------------------------------------------ tools + policy
def test_office_tools_registered_and_levels():
    t = tools.build()
    for name in ("office_create_docx", "office_create_xlsx", "office_create_pptx", "office_read", "office_edit",
                 "office_convert", "office_recalc", "office_render", "office_check", "office_review"):
        assert name in t and t[name].group == "office"
    assert "office_read" not in tools.build(office=False)


def test_office_policy(ws, tmp_path, monkeypatch):
    assert policy.classify("office_create_docx", "risky", {"path": "scratch/a.docx"})[0] == "safe"
    assert policy.classify("office_create_docx", "risky", {"path": "a.docx"})[0] == "risky"
    assert policy.classify("office_create_xlsx", "risky", {"path": "/etc/x.xlsx"})[0] == "blocked"
    assert policy.classify("office_read", "safe", {"path": "a.docx"})[0] == "safe"
    assert policy.classify("office_read", "safe", {"path": "~/.ssh/id_rsa"})[0] == "blocked"
    assert policy.classify("office_read", "safe", {"path": ".env"})[0] == "blocked"
    assert policy.classify("office_edit", "risky", {"path": "a.docx"})[0] == "risky"
    assert policy.classify("office_convert", "risky", {"path": "scratch/a.docx", "to": "pdf"})[0] == "safe"
    assert policy.classify("office_convert", "risky", {"path": "a.docx", "to": "pdf"})[0] == "risky"
    assert policy.classify("office_convert", "risky", {"path": "a.docx", "to": "pdf", "out": "scratch/a.pdf"})[0] == "safe"
    assert policy.classify("office_recalc", "safe", {"path": "a.xlsx"})[0] == "safe"
    assert policy.classify("office_recalc", "safe", {"path": "a.xlsx", "save": True})[0] == "risky"
    monkeypatch.setenv("AUTOMOTO_ALLOWED_ROOTS", str(tmp_path))
    assert policy.classify("office_read", "safe", {"path": str(tmp_path / "x.docx")})[0] == "risky"


def test_untrusted_office_outputs():
    from automoto import loop

    t = tools.build()
    for n in ("office_read", "office_check", "office_review"):
        assert loop.is_untrusted(t[n]) if hasattr(loop, "is_untrusted") else n in loop.UNTRUSTED_TOOLS


def test_tool_call_via_registry(ws):
    t = tools.build()
    ctx = tools.ToolContext()
    ok, out = tools.call(t["office_create_docx"], ctx, {"path": "scratch/t.docx", "spec": {"blocks": [{"type": "heading", "text": "Hi"}]}})
    assert ok, out
    ok, out = tools.call(t["office_read"], ctx, {"path": "scratch/t.docx"})
    assert ok and "# Hi" in out
    ok, out = tools.call(t["office_create_docx"], ctx, {"path": "scratch/t.txt", "spec": {}})
    assert not ok and ".docx" in out
    ok, out = tools.call(t["office_create_pptx"], ctx, {"path": "scratch/t.pptx", "spec": "{bad json"})
    assert not ok and "JSON" in out


def test_review_uses_vision(ws, monkeypatch):
    from automoto import office_tools, router

    img = ws / "scratch" / "renders" / "x" / "page-1.png"
    img.parent.mkdir(parents=True, exist_ok=True)
    from PIL import Image

    Image.new("RGB", (300, 420), "white").save(img)
    (ws / "x.pdf").write_bytes(b"%PDF")
    monkeypatch.setattr(office, "render", lambda *a, **k: [img])
    seen = {}

    class A:
        text, model = "looks fine", "vl-model"

    def fake_ask(messages, **kw):
        seen.update(kw)
        seen["parts"] = messages[0]["content"]
        return A()
    monkeypatch.setattr(router, "ask", fake_ask)
    ok, out = tools.call(tools.build()["office_review"], tools.ToolContext(keys={"nvidia": "k"}), {"path": "x.pdf"})
    assert ok and "looks fine" in out and seen["need_vision"] is True and seen["keys"] == {"nvidia": "k"}
    assert sum(1 for p in seen["parts"] if p.get("type") == "image_url") == 1  # one contact sheet
    assert office_tools  # imported


# ------------------------------------------------------------------ services
def test_cloud_office_services():
    from automoto import services

    for sid in ("google-docs", "google-sheets", "google-slides", "msgraph"):
        assert sid in services.SERVICES
    c = services.classify
    assert c({"service": "msgraph", "method": "POST", "path": "/v1.0/me/sendMail"})[0] == "blocked"
    assert c({"service": "msgraph", "method": "GET", "path": "/v1.0/me/messages"})[0] == "blocked"
    assert c({"service": "msgraph", "method": "POST", "path": "/v1.0/me/drive/items/1/createLink"})[0] == "blocked"
    assert c({"service": "msgraph", "method": "POST", "path": "/v1.0/me/drive/items/1/invite"})[0] == "blocked"
    assert c({"service": "msgraph", "method": "GET", "path": "/v1.0/me/drive/root/children"})[0] == "safe"
    assert c({"service": "google-docs", "method": "POST", "path": "/v1/documents/abc:batchUpdate"})[0] == "risky"
    assert c({"service": "google-sheets", "method": "GET", "path": "/v4/spreadsheets/x/values/A1:B2"})[0] == "safe"
    assert c({"service": "github", "method": "GET", "path": "/repos/a/b/events"})[0] == "safe"
    with pytest.raises(services.ServiceError, match="token"):
        services.call("google-docs", "GET", "/v1/documents/x")


# ------------------------------------------------------------------ server files API
@pytest.fixture
def client():
    from fastapi.testclient import TestClient

    from automoto import server

    return TestClient(server.app), {"X-AutoMoto-Token": server.SESSION_TOKEN}


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def test_files_upload_list_download(client, ws):
    c, H = client
    d = ws / "src.docx"
    office_build.build_docx({"blocks": [{"type": "heading", "text": "Up"}]}, d)
    r = c.post("/api/files/upload", json={"name": "../../My Report.docx", "data": _b64(d.read_bytes())}, headers=H)
    assert r.status_code == 200, r.text
    assert r.json()["path"] == "uploads/My Report.docx"
    r2 = c.post("/api/files/upload", json={"name": "My Report.docx", "data": _b64(d.read_bytes())}, headers=H)
    assert r2.json()["path"] == "uploads/My Report (2).docx"
    lst = c.get("/api/files?dir=uploads", headers=H).json()
    assert {i["name"] for i in lst["items"]} == {"My Report.docx", "My Report (2).docx"}
    dl = c.get("/api/files/download", params={"path": "uploads/My Report.docx"}, headers=H)
    assert dl.status_code == 200 and dl.content == d.read_bytes()
    assert dl.headers["content-type"] == "application/octet-stream" and "attachment" in dl.headers["content-disposition"]
    assert c.get("/api/files", headers={"X-AutoMoto-Token": "bad"}).status_code == 403


def test_files_refuses_bad_input(client, ws):
    c, H = client
    assert c.post("/api/files/upload", json={"name": ".env", "data": _b64(b"A=1")}, headers=H).status_code == 400
    assert c.post("/api/files/upload", json={"name": "x.docx", "data": "!!notbase64"}, headers=H).status_code == 400
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x/>')
    r = c.post("/api/files/upload", json={"name": "x.docx", "data": _b64(buf.getvalue())}, headers=H)
    assert r.status_code == 400 and "XXE" in r.text
    assert not list((ws / "uploads").glob("*")) if (ws / "uploads").exists() else True
    for bad in ("../../etc/passwd", "/etc/passwd", ".env"):
        assert c.get("/api/files/download", params={"path": bad}, headers=H).status_code in (400, 404)
    assert c.get("/api/files", params={"dir": "../.."}, headers=H).status_code == 400


def test_files_preview(client, ws):
    c, H = client
    from PIL import Image

    Image.new("RGB", (10, 10)).save(ws / "p.png")
    r = c.get("/api/files/preview", params={"path": "p.png"}, headers=H)
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    (ws / "x.svg").write_text("<svg onload=alert(1)/>")
    assert c.get("/api/files/preview", params={"path": "x.svg"}, headers=H).status_code == 415
    if HAS_LO and HAS_POPPLER:
        office_build.build_docx({"blocks": [{"type": "heading", "text": "Preview"}]}, ws / "pv.docx")
        r = c.get("/api/files/preview", params={"path": "pv.docx"}, headers=H)
        assert r.status_code == 200 and r.content[:4] == b"\x89PNG"


@needs_lo
def test_toc_ignores_heading_words_in_body_text(ws):
    filler = "Wind and solar generation keep growing across the country every single year. " * 30
    spec = {"properties": {"title": "T"}, "blocks": [
        {"type": "toc"},
        {"type": "heading", "level": 1, "text": "Introduction"},
        {"type": "paragraph", "text": "This report looks at the challenges of the grid. " + filler},
        {"type": "paragraph", "text": filler}, {"type": "paragraph", "text": filler},
        {"type": "heading", "level": 1, "text": "Challenges"},
        {"type": "paragraph", "text": "Grid and storage."}]}
    out = ws / "t.docx"
    office_build.create_docx(spec, out)
    pdf = office.lo_convert(out, "pdf", ws / "o")
    from pypdf import PdfReader

    pages = [p.extract_text() for p in PdfReader(str(pdf)).pages]
    real = next(i + 1 for i, t in enumerate(pages) if i > 0 and any(ln.strip() == "Challenges" for ln in t.splitlines()))
    assert real >= 3
    assert f"Challenges\t{real}" in office.read(str(out))


def test_references_heading_not_duplicated_and_lint(ws):
    out = ws / "refs.docx"
    office_build.build_docx({"properties": {"title": "T"}, "blocks": [
        {"type": "heading", "text": "Intro"}, {"type": "heading", "text": "5. References"},
        {"type": "references", "items": [{"authors": "A, B.", "year": 2020, "title": "X"}]}]}, out)
    t = office.read(str(out))
    assert t.count("References") == 1
    out2 = ws / "dup.docx"
    office_build.build_docx({"properties": {"title": "T"}, "blocks": [
        {"type": "heading", "text": "1. Results"}, {"type": "heading", "text": "Results"}]}, out2)
    assert "duplicate heading" in office_check.check(str(out2))


def test_pptx_placeholders_keep_geometry(ws):
    out = ws / "g.pptx"
    office_build.create_pptx({"slides": [{"layout": "bullets", "title": "T", "bullets": ["a", "b"], "notes": "n"}]}, out)
    import pptx

    prs = pptx.Presentation(str(out))
    body = [ph for ph in prs.slides[0].placeholders if ph.placeholder_format.idx == 1][0]
    assert body.height > 0 and body.top > 0 and body.left + body.width <= prs.slide_width
    assert "past the slide edge" not in office_check.check(str(out))


def test_chart_with_non_adjacent_columns(ws):
    out = ws / "na.xlsx"
    rows = [["M", "Rev", "Cost", "Profit"], ["Jan", 10, 4, "=B2-C2"], ["Feb", 12, 5, "=B3-C3"]]
    office_build.create_xlsx({"sheets": [{"name": "S", "rows": rows, "charts": [
        {"type": "column", "data": ["B1:B3", "D1:D3"], "categories": "A2:A3"},
        {"type": "line", "data": "$B$1:$B$3,$D$1:$D$3", "categories": "$A$2:$A$3"}]}]}, out)
    import openpyxl

    charts = openpyxl.load_workbook(out)["S"]._charts
    assert len(charts) == 2 and all(len(c.series) == 2 for c in charts)
