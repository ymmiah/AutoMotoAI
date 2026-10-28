"""Publication checks ("lint") for Office files and PDFs.

Each check returns findings as ``(severity, message)`` where severity is ``error``, ``warn``
or ``info``. The checks are heuristics: they catch the usual problems (skipped heading levels,
images without alt text, text that will overflow a slide, fonts not embedded in a PDF, formula
errors) but are not a replacement for looking at the rendered pages (``office_render``).
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

from automoto import office
from automoto.office import OfficeError

Finding = tuple[str, str]


def check_pdf(p: Path) -> list[Finding]:
    out: list[Finding] = []
    info = office.pdf_info(p)
    if info:
        out.append(("info", f"{info.get('Pages', '?')} page(s), size {info.get('Page size', '?')}"))
        if info.get("Tagged", "no").lower() != "yes":
            out.append(("warn", "PDF is not tagged - screen readers can't follow its structure (export with UseTaggedPDF)"))
        if not info.get("Title"):
            out.append(("warn", "PDF has no Title metadata - set the document title before exporting"))
        if info.get("Encrypted", "no").lower().startswith("yes"):
            out.append(("warn", "PDF is encrypted"))
    fonts = office.pdf_fonts(p)
    bad = sorted({f["name"] for f in fonts if f["emb"] != "yes"})
    if bad:
        out.append(("error", f"fonts not embedded (printers/publishers will reject this): {', '.join(bad[:10])}"))
    elif fonts:
        out.append(("info", f"all {len(fonts)} font(s) embedded"))
    try:
        from pypdf import PdfReader

        r = PdfReader(str(p))
        if r.outline == [] and len(r.pages) > 4:
            out.append(("warn", "no bookmarks/outline - long documents should have them (use real heading styles)"))
        text = "".join((pg.extract_text() or "") for pg in r.pages[:60])
        if "Error! Reference source not found" in text or "Error! Bookmark not defined" in text:
            out.append(("error", "broken cross-reference text found in the PDF"))
        if len(r.pages) and not text.strip():
            out.append(("warn", "no extractable text - is this a scanned image? It won't be searchable or accessible"))
    except Exception as exc:  # pragma: no cover - damaged PDFs
        out.append(("warn", f"could not parse the PDF fully: {exc}"))
    return out


def check_docx(p: Path) -> list[Finding]:
    import docx
    from docx.oxml.ns import qn

    out: list[Finding] = []
    d = docx.Document(str(p))
    levels = []
    for para in d.paragraphs:
        name = para.style.name if para.style is not None else ""
        m = re.match(r"Heading (\d)", name)
        if m and para.text.strip():
            levels.append((int(m.group(1)), para.text.strip()))
        elif name.startswith("Heading") and not para.text.strip():
            out.append(("warn", "empty heading paragraph"))
    if not levels:
        out.append(("warn", "no headings - use Heading 1/2/3 styles so the TOC, PDF bookmarks and screen readers work"))
    else:
        if levels[0][0] != 1:
            out.append(("warn", f"first heading is level {levels[0][0]} ('{levels[0][1][:40]}'), expected level 1"))
        for (a, _), (b, t) in zip(levels, levels[1:]):
            if b > a + 1:
                out.append(("warn", f"heading level jumps from {a} to {b} at '{t[:50]}'"))
        def bare(t: str) -> str:
            return re.sub(r"^[\d.\s]+", "", t).strip().lower()
        for (_, t1), (_, t2) in zip(levels, levels[1:]):
            if bare(t1) and bare(t1) == bare(t2):
                out.append(("warn", f"duplicate heading '{t2[:50]}' right after '{t1[:50]}'"))
    body = d.element.body
    for dp in body.iter(qn("wp:docPr")):
        if not (dp.get("descr") or "").strip():
            out.append(("warn", f"image '{dp.get('name', '?')}' has no alt text"))
    xml = body.xml
    if "TOC \\o" in xml or 'TOC \\\\o' in xml:
        toc_tabs = re.findall(r"<w:t[^>]*>\t?0</w:t>", xml)
        if toc_tabs:
            out.append(("warn", "table of contents still shows placeholder page numbers - update fields (F9) or rebuild with office_create_docx"))
    bookmarks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    for instr in body.iter(qn("w:instrText")):
        m = re.match(r"\s*(?:REF|PAGEREF)\s+(\S+)", instr.text or "")
        if m and m.group(1) not in bookmarks:
            out.append(("error", f"cross-reference to missing bookmark '{m.group(1)}'"))
    if not (d.core_properties.title or "").strip():
        out.append(("warn", "no document title in properties (used as the PDF title)"))
    info = office.inspect_file(p)
    if info["external_links"]:
        out.append(("info", f"{len(info['external_links'])} external link(s): {', '.join(info['external_links'][:5])}"))
    empty = sum(1 for para in d.paragraphs if not para.text.strip() and not para.runs)
    if empty > 25:
        out.append(("info", f"{empty} empty paragraphs - use spacing/page breaks instead of blank lines"))
    out.append(("info", f"{len(levels)} heading(s), {len(d.tables)} table(s), {len(d.inline_shapes)} inline image(s)"))
    return out


def check_xlsx(p: Path) -> list[Finding]:
    import openpyxl

    out: list[Finding] = []
    wb = openpyxl.load_workbook(str(p), keep_links=False)
    names = set(wb.sheetnames)
    nform = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    nform += 1
                    for ref in re.findall(r"(?:'([^']+)'|([A-Za-z_][\w.]*))!", v):
                        sh = ref[0] or ref[1]
                        if sh not in names:
                            out.append(("error", f"{ws.title}!{c.coordinate} refers to missing sheet '{sh}'"))
                    if "#REF!" in v:
                        out.append(("error", f"{ws.title}!{c.coordinate} formula contains #REF!"))
                    if re.search(r"\[\d+\]", v):
                        out.append(("warn", f"{ws.title}!{c.coordinate} uses an external workbook link"))
        if ws.max_row > 30 and ws.freeze_panes is None:
            out.append(("info", f"sheet '{ws.title}' has {ws.max_row} rows and no frozen header row"))
    out.append(("info", f"{len(names)} sheet(s), {nform} formula(s)"))
    if nform and office.lo_available():
        res = office.recalc(str(p))
        m = re.search(r"; (\d+) error", res)
        if m and int(m.group(1)):
            errs = res.split("ERRORS - fix these formulas:\n", 1)[-1].split("\n")[:10]
            out.append(("error", f"{m.group(1)} formula error(s) after recalculation: " + "; ".join(e for e in errs if "!" in e)))
        else:
            out.append(("info", "all formulas recalculated without errors (LibreOffice)"))
    elif nform:
        out.append(("info", "LibreOffice not installed - formulas were not recalculated"))
    return out


def _est_lines(text: str, width_pt: float, size_pt: float) -> int:
    per_line = max(1, int(width_pt / (0.5 * size_pt)))
    return sum(max(1, -(-len(line) // per_line)) for line in text.split("\n"))


def check_pptx(p: Path) -> list[Finding]:
    import pptx
    from pptx.util import Pt

    out: list[Finding] = []
    prs = pptx.Presentation(str(p))
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, 1):
        title = s.shapes.title
        if title is None or not title.text_frame.text.strip():
            kinds = {sh.shape_type for sh in s.shapes}
            if len(s.shapes) > 1 and i != 1:
                out.append(("info", f"slide {i}: no title (screen readers and the outline use slide titles)"))
        for sh in s.shapes:
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
                out.append(("warn", f"slide {i}: empty placeholder '{sh.name}' (shows 'Click to add text' in PowerPoint)"))
            if sh.left is not None and sh.width is not None and (
                    sh.left < -Pt(2) or sh.top < -Pt(2) or sh.left + sh.width > W + Pt(2) or sh.top + sh.height > H + Pt(2)):
                out.append(("warn", f"slide {i}: '{sh.name}' extends past the slide edge"))
            if sh.shape_type == 13:  # picture
                descr = sh._element.nvPicPr.cNvPr.get("descr")
                if not (descr or "").strip():
                    out.append(("warn", f"slide {i}: picture '{sh.name}' has no alt text"))
            if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.width is None:
                continue
            tf = sh.text_frame
            sizes = [r.font.size.pt for para in tf.paragraphs for r in para.runs if r.font.size]
            size = max(set(sizes), key=sizes.count) if sizes else 18.0
            if sizes and min(sizes) < 12 and sh.name not in ("slide number", "footer"):
                out.append(("warn", f"slide {i}: text under 12 pt in '{sh.name}' - hard to read when projected"))
            width_pt = (sh.width - Pt(14)) / 12700
            height_pt = sh.height / 12700
            lines = _est_lines(tf.text, width_pt, size)
            need = lines * size * 1.25
            if need > height_pt * 1.08 and sh.name not in ("slide number", "footer"):
                out.append(("warn", f"slide {i}: text in '{sh.name}' probably overflows (~{lines} lines at {size:.0f} pt "
                                    f"need {need:.0f} pt, box is {height_pt:.0f} pt) - shorten it or split the slide"))
            words = len(tf.text.split())
            if words > 90:
                out.append(("warn", f"slide {i}: {words} words - aim for one idea per slide and move detail to the notes"))
        if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
            out.append(("info", f"slide {i}: no speaker notes"))
    out.append(("info", f"{len(prs.slides)} slide(s), {W / 914400:.2f}×{H / 914400:.2f} in"))
    return out


def check(path: str) -> str:
    p = office.resolve_in(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        findings = check_pdf(p)
    elif ext in (".docx", ".docm", ".dotx"):
        findings = check_docx(p)
    elif ext in (".xlsx", ".xlsm", ".xltx"):
        findings = check_xlsx(p)
    elif ext in (".pptx", ".pptm", ".potx"):
        findings = check_pptx(p)
    elif ext in office.LEGACY | office.ODF:
        target = {"writer": "docx", "calc": "xlsx", "impress": "pptx"}[office._family(ext)]
        with tempfile.TemporaryDirectory(prefix="automoto-check-") as tmp:
            conv = office.lo_convert(p, target, Path(tmp))
            findings = [("info", f"converted from {ext} to check it")] + {"docx": check_docx, "xlsx": check_xlsx,
                                                                            "pptx": check_pptx}[target](conv)
    else:
        raise OfficeError(f"can't check {ext} files")
    info = office.inspect_file(p) if ext != ".pdf" else {"macros": False}
    if info.get("macros"):
        findings.insert(0, ("warn", "file contains macros (they are never run by AutoMoto)"))
    order = {"error": 0, "warn": 1, "info": 2}
    findings.sort(key=lambda f: order[f[0]])
    n_err = sum(1 for f in findings if f[0] == "error")
    n_warn = sum(1 for f in findings if f[0] == "warn")
    icon = {"error": "✗", "warn": "⚠", "info": "·"}
    head = f"{p.name}: {n_err} error(s), {n_warn} warning(s)" + (" - looks publication-ready" if not n_err and not n_warn else "")
    return head + "\n" + "\n".join(f"{icon[s]} {m}" for s, m in findings[:120])
