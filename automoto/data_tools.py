"""Agent tools for data files (group ``data``)."""

from __future__ import annotations

import json
from typing import Any

from automoto import datafiles as dfm
from automoto import sandbox
from automoto.archive_tools import ARR, B, _list
from automoto.tools import I, S, Tool, ToolContext, ToolFailed, _obj

OBJ = {"type": "object"}


def _wrap(fn):
    def inner(ctx: ToolContext, **kw: Any) -> str:
        try:
            return fn(ctx, **kw)
        except (dfm.DataError, sandbox.SandboxError) as exc:
            raise ToolFailed(str(exc)) from None
        except UnicodeDecodeError as exc:
            raise ToolFailed(f"could not decode the file ({exc.encoding}); pass encoding='cp1252' or similar") from None
    return inner


def _json_arg(v: Any, what: str) -> Any:
    if isinstance(v, str) and v.strip()[:1] in "[{":
        try:
            return json.loads(v)
        except json.JSONDecodeError as exc:
            raise ToolFailed(f"{what} is not valid JSON: {exc}") from None
    return v


@_wrap
def data_inspect(ctx: ToolContext, path: str, sheet: str = "", table: str = "", rows: int = 8) -> str:
    p = dfm.resolve_in(path)
    k = dfm.kind_of(p)
    size = p.stat().st_size
    head = f"{sandbox.rel(p)}: {k}, {size:,} bytes"
    n = max(1, min(int(rows or 8), 50))
    if k in ("csv", "excel", "jsonl", "parquet") or (k == "json" and size < dfm.MAX_LOAD()):
        if k == "json":
            obj = dfm.load_structured(p)
            if not (isinstance(obj, list) or (isinstance(obj, dict) and any(isinstance(v, list) for v in obj.values()))):
                return head + "\n" + "\n".join(dfm.outline(obj))
        extra = ""
        if k == "csv":
            d = dfm.sniff_csv(p)
            lines, more = dfm.count_lines(p, dfm.MAX_ROWS())
            extra = (f"\nencoding {d['encoding']}, delimiter {d['delimiter']!r}, header {'yes' if d['header'] else 'no'}, "
                     f"~{lines:,}{'+' if more else ''} lines")
        df, info = dfm.load_table(p, sheet=sheet or None, nrows=200_000, table=table or None)
        trunc = " (first 200,000 rows profiled)" if info.get("truncated_at_rows") else ""
        if k == "excel":
            from openpyxl import load_workbook

            try:
                wb = load_workbook(p, read_only=True)
                extra = f"\nsheets: {', '.join(wb.sheetnames)}"
                wb.close()
            except Exception:
                pass
        return f"{head}{extra}{trunc}\n{dfm.describe(df)}\n\npreview:\n{dfm.to_markdown(df, n)}"
    if k in ("json", "yaml", "toml", "xml"):
        obj = dfm.load_structured(p)
        return head + "\n" + "\n".join(dfm.outline(obj, max_depth=4))
    if k == "sqlite":
        sch = dfm.sqlite_schema(p)
        return head + "\n" + "\n".join(f"- {t['type']} {t['name']} ({t['rows']} rows): {', '.join(t['columns'][:30])}"
                                       for t in sch)
    if k == "image":
        return head + "\n" + json.dumps(dfm.image_info(p), ensure_ascii=False)
    if k == "pdf":
        r = dfm.pdf_text(p, 1, 1, max_chars=1500)
        return f"{head}\n{r['pages']} page(s), metadata {json.dumps(r['metadata'], ensure_ascii=False)}\npage 1 text:\n{r['text']}"
    if k == "markdown":
        text = p.read_text(encoding=dfm.detect_encoding(p), errors="replace")
        ol = dfm.markdown_outline(text)
        return f"{head}, {len(text.splitlines()):,} lines, {len(text.split()):,} words\noutline:\n" + "\n".join(ol[:80])
    if k == "zip":
        return head + "\nthis is an archive - use archive_inspect / archive_list"
    s = dfm.log_summary(p)
    return (f"{head}, {s['lines']:,} lines" + (f", levels {s['levels']}" if s["levels"] else "")
            + (f", {s['first_timestamp']} → {s['last_timestamp']}" if s["first_timestamp"] else "")
            + f"\nlast lines:\n{dfm.tail(p, n)}")


@_wrap
def data_query(ctx: ToolContext, path: str, columns: Any = None, filters: Any = None, group_by: Any = None,
               aggregate: Any = None, sort: Any = None, limit: int = 0, dedupe: bool = False, out: str = "",
               sheet: str = "", table: str = "", delimiter: str = "", encoding: str = "", show: int = 25) -> str:
    df, info = dfm.load_table(path, sheet=sheet or None, table=table or None, delimiter=delimiter or None,
                              encoding=encoding or None)
    filters = _json_arg(filters, "filters")
    if isinstance(filters, dict):
        filters = [filters]
    aggregate = _json_arg(aggregate, "aggregate")
    sort = _json_arg(sort, "sort")
    res = dfm.query(df, columns=_list(columns), filters=filters, group_by=_list(group_by),
                    aggregate=aggregate if isinstance(aggregate, dict) else None,
                    sort=sort if isinstance(sort, list) else _list(sort), limit=int(limit or 0) or None,
                    dedupe=bool(dedupe))
    text = f"{len(df):,} rows in → {len(res):,} rows out" + \
        (f" (input truncated at {info['truncated_at_rows']:,} rows)" if info.get("truncated_at_rows") else "") + "\n"
    text += dfm.to_markdown(res, max(1, min(int(show or 25), 200)))
    if out:
        w = dfm.write_table(res, out)
        text += f"\nwrote {w['out']} ({w['rows']:,} rows, {w['bytes']:,} bytes)"
    return text


@_wrap
def data_convert(ctx: ToolContext, path: str, out: str, sheet: str = "", table: str = "", select: str = "",
                 delimiter: str = "", encoding: str = "") -> str:
    p = dfm.resolve_in(path)
    k = dfm.kind_of(p)
    o_ext = out.lower().rsplit(".", 1)[-1]
    structured_out = o_ext in ("yaml", "yml", "toml") or (o_ext in ("json", "jsonl") and k in ("yaml", "toml", "xml", "json"))
    if k in ("yaml", "toml", "xml") or (structured_out and k in ("json", "jsonl")):
        obj = dfm.select(dfm.load_structured(p), select) if select else dfm.load_structured(p)
        if o_ext in ("csv", "tsv", "xlsx", "parquet", "md", "html"):
            pd = dfm._pd()
            rows = obj if isinstance(obj, list) else [obj]
            w = dfm.write_table(pd.json_normalize(rows, max_level=1), out)
        else:
            w = dfm.write_structured(obj, out)
        return f"converted {sandbox.rel(p)} → {w['out']} ({w['bytes']:,} bytes)"
    df, info = dfm.load_table(p, sheet=sheet or None, table=table or None, delimiter=delimiter or None,
                              encoding=encoding or None)
    w = dfm.write_table(df, out, sheet=sheet or "Data")
    note = f" (input truncated at {info['truncated_at_rows']:,} rows)" if info.get("truncated_at_rows") else ""
    return f"converted {sandbox.rel(p)} → {w['out']} ({w['rows']:,} rows x {w['columns']} columns, {w['bytes']:,} bytes){note}"


@_wrap
def data_read(ctx: ToolContext, path: str, select: str = "", max_chars: int = 15000) -> str:
    p = dfm.resolve_in(path)
    obj = dfm.load_structured(p)
    if select:
        obj = dfm.select(obj, select)
    text = json.dumps(obj, indent=1, ensure_ascii=False, default=str)
    cap = max(500, min(int(max_chars or 15000), 60000))
    return text if len(text) <= cap else text[:cap] + f"\n[… {len(text) - cap} more chars - use select='a.b[0]' to narrow]"


@_wrap
def sqlite_query(ctx: ToolContext, path: str, sql: str = "", params: Any = None, limit: int = 200, out: str = "") -> str:
    p = dfm.resolve_in(path)
    if not sql.strip():
        sch = dfm.sqlite_schema(p)
        return "\n".join(f"- {t['type']} {t['name']} ({t['rows']} rows): {', '.join(t['columns'][:40])}" for t in sch)
    cols, rows, more = dfm.sqlite_query(p, sql, params=_json_arg(params, "params") or None,
                                        limit=max(1, min(int(limit or 200), 100000)))
    pd = dfm._pd()
    df = pd.DataFrame(rows, columns=cols)
    text = f"{len(rows):,} row(s){' (more available - raise limit)' if more else ''}\n" + dfm.to_markdown(df, 50)
    if out:
        w = dfm.write_table(df, out)
        text += f"\nwrote {w['out']}"
    return text


@_wrap
def text_search(ctx: ToolContext, path: str, pattern: str, regex: bool = False, ignore_case: bool = True,
                context: int = 0, max_matches: int = 100, invert: bool = False) -> str:
    out, total, lines = dfm.grep(path, pattern, regex=bool(regex), ignore_case=ignore_case is not False,
                                 context=int(context or 0), max_matches=max(1, min(int(max_matches or 100), 1000)),
                                 invert=bool(invert))
    return f"{total:,} matching line(s) of {lines:,}\n" + "\n".join(out)


@_wrap
def text_tail(ctx: ToolContext, path: str, lines: int = 50) -> str:
    return dfm.tail(path, lines)


@_wrap
def pdf_text(ctx: ToolContext, path: str, first: int = 1, last: int = 0, layout: bool = False, max_chars: int = 30000) -> str:
    r = dfm.pdf_text(path, first, last, bool(layout), max(1000, min(int(max_chars or 30000), 60000)))
    return f"{r['pages']} page(s); text of pages {r['range'][0]}-{r['range'][1]} via {r['engine']}:\n{r['text']}"


@_wrap
def image_info(ctx: ToolContext, path: str) -> str:
    return json.dumps(dfm.image_info(path), ensure_ascii=False, indent=1)


FILTERS = {"type": "array", "items": OBJ, "description": "[{column, op, value}] op: == != > >= < <= contains not_contains "
           "startswith endswith in not_in between isnull notnull regex"}

TOOLS: list[Tool] = [
    Tool("data_inspect", "Profile any data file before working on it: CSV/TSV (encoding, delimiter, rows, column types, "
         "nulls, top values, preview), Excel, JSON/JSONL, YAML, TOML, XML, Parquet, SQLite (tables, row counts), "
         "Markdown (outline), logs (levels, time range, last lines), images (size, EXIF, GPS warning), PDF (pages, metadata).",
         _obj({"path": S, "sheet": S, "table": S, "rows": I}, ["path"]), "safe", data_inspect, "data"),
    Tool("data_query", "Filter, select, group, aggregate, sort and de-duplicate a table (CSV/TSV/xlsx/JSON/JSONL/Parquet/"
         "SQLite table) with pandas; optional out= writes the result (csv, tsv, xlsx, json, jsonl, parquet, md, html, sqlite). "
         "aggregate={'Revenue':'sum','*':'count'} or {'Price':['mean','max']}; sort=['-Revenue','Region'].",
         _obj({"path": S, "columns": ARR, "filters": FILTERS, "group_by": ARR, "aggregate": OBJ,
               "sort": ARR, "limit": I, "dedupe": B, "out": S, "sheet": S, "table": S, "delimiter": S,
               "encoding": S, "show": I}, ["path"]), "safe", data_query, "data"),
    Tool("data_convert", "Convert between csv, tsv, xlsx, json, jsonl, parquet, md, html, sqlite (tables) and json, yaml, "
         "toml, xml→json (structured). select='a.b' picks part of a structured file. Formula-looking text from CSVs "
         "stays text in xlsx.", _obj({"path": S, "out": S, "sheet": S, "table": S, "select": S, "delimiter": S,
                                       "encoding": S}, ["path", "out"]), "risky", data_convert, "data"),
    Tool("data_read", "Read JSON/JSONL/YAML/TOML/XML as JSON (YAML safe_load; XML with DTDs/entities refused). "
         "select='items[0].name' or 'items.*.id' narrows it.",
         _obj({"path": S, "select": S, "max_chars": I}, ["path"]), "safe", data_read, "data"),
    Tool("sqlite_query", "Read-only SQL on an SQLite database file (SELECT/WITH/PRAGMA table_info only; writes, ATTACH "
         "and extensions are refused). No sql = list tables. out= exports the result.",
         _obj({"path": S, "sql": S, "params": {"type": "array"}, "limit": I, "out": S}, ["path"]), "safe",
         sqlite_query, "data"),
    Tool("text_search", "grep a large text/log/CSV file (streams; literal by default, regex=true for patterns) with "
         "line numbers and optional context lines.",
         _obj({"path": S, "pattern": S, "regex": B, "ignore_case": B, "context": I, "max_matches": I, "invert": B},
              ["path", "pattern"]), "safe", text_search, "data"),
    Tool("text_tail", "Last N lines of a (large) text or log file.", _obj({"path": S, "lines": I}, ["path"]), "safe",
         text_tail, "data"),
    Tool("pdf_text", "Extract the text of a PDF (page range; layout=true keeps columns) with pdftotext or pypdf.",
         _obj({"path": S, "first": I, "last": I, "layout": B, "max_chars": I}, ["path"]), "safe", pdf_text, "data"),
    Tool("image_info", "Image metadata: format, size, mode, DPI, frames, EXIF camera/date, and whether GPS location is embedded.",
         _obj({"path": S}, ["path"]), "safe", image_info, "data"),
]
