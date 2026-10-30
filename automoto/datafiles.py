"""Data files: CSV/TSV, JSON/JSONL, YAML, TOML, XML, Parquet, SQLite, Markdown, logs, images, PDF text.

Everything reads untrusted input defensively: size and row limits, streaming for counts and
log search, ``yaml.safe_load`` only, ``defusedxml`` for XML, SQLite opened read-only with an
authorizer that allows nothing but reads, Pillow's decompression-bomb guard on images, and
formula-looking strings written to .xlsx as text (CSV injection).
"""

from __future__ import annotations

import codecs
import csv
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import time
from collections import deque
from pathlib import Path
from typing import Any, Iterable

from automoto import sandbox


class DataError(ValueError):
    pass


def _env_int(name: str, default: int) -> int:
    try:
        return max(1, int(os.environ.get(name, "") or default))
    except ValueError:
        return default


MAX_FILE = lambda: _env_int("AUTOMOTO_DATA_MAX_BYTES", 1024 * 1024 * 1024)  # noqa: E731 - 1 GB to stream
MAX_LOAD = lambda: _env_int("AUTOMOTO_DATA_MAX_LOAD_BYTES", 300 * 1024 * 1024)  # noqa: E731 - into memory
MAX_ROWS = lambda: _env_int("AUTOMOTO_DATA_MAX_ROWS", 2_000_000)  # noqa: E731

TABULAR = {".csv", ".tsv", ".tab", ".psv", ".txt", ".xlsx", ".xlsm", ".xls", ".ods", ".json", ".jsonl", ".ndjson",
           ".parquet", ".pq"}
KINDS = {".csv": "csv", ".tsv": "csv", ".tab": "csv", ".psv": "csv", ".json": "json", ".jsonl": "jsonl",
         ".ndjson": "jsonl", ".geojson": "json", ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".xml": "xml",
         ".svg": "xml", ".rss": "xml", ".atom": "xml", ".kml": "xml", ".gpx": "xml", ".plist": "xml",
         ".parquet": "parquet", ".pq": "parquet", ".sqlite": "sqlite", ".sqlite3": "sqlite", ".db": "sqlite",
         ".md": "markdown", ".markdown": "markdown", ".log": "text", ".txt": "text", ".ini": "text", ".cfg": "text",
         ".conf": "text", ".env.example": "text", ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image",
         ".webp": "image", ".tif": "image", ".tiff": "image", ".bmp": "image", ".heic": "image", ".ico": "image",
         ".pdf": "pdf", ".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".ods": "excel"}


def resolve_in(path: str | os.PathLike) -> Path:
    p = sandbox.resolve(str(path), allow_outside=True)
    if not p.is_file():
        raise DataError(f"no such file: {path}")
    if p.stat().st_size > MAX_FILE():
        raise DataError(f"{p.name} is larger than the {MAX_FILE() // 2**20} MB limit")
    return p


def resolve_out(path: str | os.PathLike) -> Path:
    p = sandbox.resolve(str(path), allow_outside=True)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def kind_of(p: Path) -> str:
    with p.open("rb") as fh:
        head = fh.read(16)
    if head.startswith(b"SQLite format 3\x00"):
        return "sqlite"
    if head.startswith(b"PAR1"):
        return "parquet"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head[:8] == b"\x89PNG\r\n\x1a\n" or head[:3] == b"\xff\xd8\xff" or head[:4] in (b"GIF8", b"RIFF"):
        return "image"
    k = KINDS.get(p.suffix.lower())
    if k:
        return k
    if head[:2] == b"PK":
        return "zip"
    return "text"


# -- encoding / CSV dialect --------------------------------------------------------------------
def detect_encoding(p: Path, sample: int = 1024 * 1024) -> str:
    with p.open("rb") as fh:
        raw = fh.read(sample)
    for bom, enc in ((codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")):
        if raw.startswith(bom):
            return enc
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as exc:
        if exc.start > len(raw) - 4:  # a multibyte char cut by the sample boundary
            return "utf-8"
    # Western European text (Excel's "CSV" on Windows) is by far the most common non-UTF-8 case:
    # accept cp1252 when its non-ASCII characters look like letters, currency or punctuation.
    import unicodedata

    try:
        text = raw.decode("cp1252")
        odd = [c for c in text if ord(c) > 127]
        good = sum(unicodedata.category(c)[0] in "LSP" or c in "\u00a0" for c in odd)
        if not odd or good / len(odd) > 0.95:
            return "cp1252"
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(raw).best()
        if best and best.encoding:
            return best.encoding
    except ImportError:
        pass
    return "latin-1"


def sniff_csv(p: Path, encoding: str | None = None) -> dict:
    enc = encoding or detect_encoding(p)
    with p.open("r", encoding=enc, errors="replace", newline="") as fh:
        sample = fh.read(64 * 1024)
    delim = None
    if p.suffix.lower() in (".tsv", ".tab"):
        delim = "\t"
    elif p.suffix.lower() == ".psv":
        delim = "|"
    if not delim:
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            lines = sample.splitlines()[:20]
            counts = {d: sum(ln.count(d) for ln in lines) for d in (",", ";", "\t", "|")}
            delim = max(counts, key=counts.get) if any(counts.values()) else ","
    # header unless the first row looks like data (has numbers) and the sniffer agrees
    first = next(csv.reader(io.StringIO(sample), delimiter=delim), [])
    numeric = any(re.fullmatch(r"\s*[-+]?\d+([.,]\d+)?\s*", c or "") for c in first)
    header = True
    if numeric:
        try:
            header = csv.Sniffer().has_header(sample)
        except csv.Error:
            header = False
    decimal = "."
    if delim in (";", "\t", "|") and re.search(r"(^|[;\t|])\s*-?\d{1,3}(\.\d{3})*,\d+\s*(?=[;\t|]|$)", sample, re.M):
        decimal = ","  # European numbers: 1.234,5
    return {"encoding": enc, "delimiter": delim, "header": header, "decimal": decimal}


def count_lines(p: Path, limit: int | None = None) -> tuple[int, bool]:
    n = 0
    with p.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            n += chunk.count(b"\n")
            if limit and n > limit:
                return n, True
    return n, False


# -- loading tables --------------------------------------------------------------------------------
def _pd():
    try:
        import pandas as pd

        return pd
    except ImportError:
        raise DataError("pandas is needed for table queries - pip install pandas") from None


def load_table(path: str | os.PathLike, *, sheet: str | None = None, encoding: str | None = None,
               delimiter: str | None = None, nrows: int | None = None, table: str | None = None):
    """A pandas DataFrame from csv/tsv/xlsx/ods/json/jsonl/parquet/sqlite (table=…). Returns (df, info)."""
    pd = _pd()
    p = resolve_in(path)
    k = kind_of(p)
    info: dict = {"file": sandbox.rel(p), "kind": k}
    cap = min(int(nrows), MAX_ROWS()) if nrows else MAX_ROWS()
    if k not in ("csv", "text", "jsonl", "sqlite") and p.stat().st_size > MAX_LOAD():
        raise DataError(f"{p.name} is too big to load into memory ({p.stat().st_size // 2**20} MB > "
                        f"{MAX_LOAD() // 2**20} MB); convert it to CSV first or raise AUTOMOTO_DATA_MAX_LOAD_BYTES")
    if k in ("csv", "text"):
        d = sniff_csv(p, encoding)
        if delimiter:
            d["delimiter"] = "\t" if delimiter in ("\\t", "tab") else delimiter
        info.update(d)
        df = pd.read_csv(p, sep=d["delimiter"], encoding=d["encoding"], nrows=cap + 1, encoding_errors="replace",
                         header=0 if d["header"] else None, on_bad_lines="skip", low_memory=False,
                         decimal=d["decimal"], thousands="." if d["decimal"] == "," else None)
        df = parse_dates(df, info)
        if not d["header"]:
            df.columns = [f"col{i + 1}" for i in range(len(df.columns))]
    elif k == "excel":
        df = pd.read_excel(p, sheet_name=sheet or 0, nrows=cap + 1)
        info["sheet"] = sheet or "(first)"
    elif k == "json":
        with p.open("r", encoding=detect_encoding(p)) as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            # {"data": [...]} / {"items": [...]} / GeoJSON features
            lst = next((v for key, v in data.items() if isinstance(v, list) and v and isinstance(v[0], dict)), None)
            if lst is None:
                data = [data]
            else:
                info["records_key"] = next(key for key, v in data.items() if v is lst)
                data = lst
        if not isinstance(data, list):
            raise DataError("JSON is not a list of records")
        df = pd.json_normalize(data[: cap + 1], max_level=1)
    elif k == "jsonl":
        rows = []
        with p.open("r", encoding=detect_encoding(p), errors="replace") as fh:
            for i, line in enumerate(fh):
                if len(rows) > cap:
                    break
                line = line.strip()
                if not line:
                    continue
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError:
                    info.setdefault("bad_lines", []).append(i + 1)
        df = pd.json_normalize(rows, max_level=1)
        if info.get("bad_lines"):
            info["bad_lines"] = info["bad_lines"][:20]
    elif k == "parquet":
        try:
            df = pd.read_parquet(p)
        except ImportError:
            raise DataError("Parquet needs pyarrow - pip install pyarrow (optional, ~150 MB)") from None
        df = df.head(cap + 1)
    elif k == "sqlite":
        con = sqlite_connect(p)
        try:
            tables = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type IN ('table','view') ORDER BY name")]
            t = table or (tables[0] if len(tables) == 1 else None)
            if not t or t not in tables:
                raise DataError(f"give table=… (one of {', '.join(tables[:20])})")
            df = pd.read_sql_query(f'SELECT * FROM "{t}" LIMIT {cap + 1}', con)  # nosec B608 - name checked against sqlite_master
            info["table"] = t
        finally:
            con.close()
    else:
        raise DataError(f"{p.name}: {k} is not a table format (csv, tsv, xlsx, json, jsonl, parquet, sqlite)")
    if len(df) > cap:
        df = df.head(cap)
        info["truncated_at_rows"] = cap
    return df, info


_DATE = re.compile(r"^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?\s*$")
_ISO = re.compile(r"^\s*\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2})?(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?\s*$")


def parse_dates(df, info: dict | None = None):
    """Turn text columns that are clearly dates into datetimes (dd/mm/yyyy unless the data says mm/dd)."""
    pd = _pd()
    done = {}
    for c in df.columns:
        s = df[c]
        if not (s.dtype == object or pd.api.types.is_string_dtype(s)):
            continue
        sample = s.dropna().astype(str).head(500)
        if sample.empty:
            continue
        if sample.map(lambda v: bool(_ISO.match(v))).mean() > 0.9:
            df[c] = pd.to_datetime(s, errors="coerce", utc=False, format="mixed")
            done[str(c)] = "ISO"
            continue
        m = sample.map(lambda v: _DATE.match(v))
        if m.map(bool).mean() < 0.9:
            continue
        firsts = [int(x.group(1)) for x in m if x]
        seconds = [int(x.group(2)) for x in m if x]
        dayfirst = not (max(seconds, default=0) > 12 and max(firsts, default=0) <= 12)
        df[c] = pd.to_datetime(s, errors="coerce", dayfirst=dayfirst, format="mixed")
        done[str(c)] = "dd/mm/yyyy" if dayfirst else "mm/dd/yyyy"
    if info is not None and done:
        info["dates_parsed"] = done
    return df


# -- derived columns (a tiny safe expression language, no eval) ------------------------------------
_FUNCS = {"month", "year", "day", "quarter", "weekday", "yearmonth", "date", "round", "abs", "upper", "lower",
          "strip", "len", "int", "float", "str"}


def derive(df, columns: dict):
    """{'Revenue': '[Units] * [Unit Price]', 'Month': 'yearmonth([Order Date])'} - columns in [brackets] or bare
    names; + - * / // % **, numbers, strings and the functions above. Parsed with ``ast``; nothing is executed."""
    import ast

    pd = _pd()
    if not isinstance(columns, dict):
        raise DataError("add_columns must be {name: expression}")
    for name, expr in columns.items():
        expr = str(expr)
        if len(expr) > 300:
            raise DataError("expression too long")
        refs: dict[str, str] = {}

        def sub(m):
            key = f"__c{len(refs)}"
            refs[key] = _col(df, m.group(1))
            return key
        code = re.sub(r"\[([^\]]+)\]", sub, expr)
        try:
            tree = ast.parse(code, mode="eval")
        except SyntaxError as exc:
            raise DataError(f"bad expression {expr!r}: {exc.msg}") from None

        def ev(node):
            if isinstance(node, ast.Expression):
                return ev(node.body)
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float, str)):
                return node.value
            if isinstance(node, ast.Name):
                if node.id in refs:
                    return df[refs[node.id]]
                return df[_col(df, node.id)]
            if isinstance(node, ast.BinOp):
                a, b = ev(node.left), ev(node.right)
                ops = {ast.Add: lambda: a + b, ast.Sub: lambda: a - b, ast.Mult: lambda: a * b, ast.Div: lambda: a / b,
                       ast.FloorDiv: lambda: a // b, ast.Mod: lambda: a % b, ast.Pow: lambda: a ** b}
                f = ops.get(type(node.op))
                if not f:
                    raise DataError("operator not allowed")
                try:
                    return f()
                except TypeError as exc:
                    raise DataError(f"type error in {expr!r}: {exc} (numbers stored as text? check the delimiter/decimal)") from None
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
                v = ev(node.operand)
                return -v if isinstance(node.op, ast.USub) else v
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in _FUNCS and not node.keywords:
                args = [ev(a) for a in node.args]
                fn = node.func.id
                x = args[0] if args else None
                if fn in ("month", "year", "day", "quarter", "weekday", "yearmonth", "date"):
                    x = pd.to_datetime(x, errors="coerce", dayfirst=True)
                    return {"month": lambda: x.dt.month, "year": lambda: x.dt.year, "day": lambda: x.dt.day,
                            "quarter": lambda: x.dt.quarter, "weekday": lambda: x.dt.day_name(),
                            "yearmonth": lambda: x.dt.strftime("%Y-%m"), "date": lambda: x.dt.date}[fn]()
                if fn == "round":
                    return x.round(int(args[1]) if len(args) > 1 else 0) if hasattr(x, "round") else round(x, int(args[1]) if len(args) > 1 else 0)
                if fn == "abs":
                    return x.abs() if hasattr(x, "abs") else abs(x)
                if fn in ("upper", "lower", "strip", "len"):
                    st = x.astype("string").str
                    return getattr(st, fn)()
                if fn in ("int", "float"):
                    return pd.to_numeric(x, errors="coerce")
                if fn == "str":
                    return x.astype("string")
            raise DataError(f"not allowed in an expression: {ast.dump(node)[:60]}")
        df = df.assign(**{str(name): ev(tree)})
    return df


# -- queries ---------------------------------------------------------------------------------------
_OPS = {"==", "!=", ">", ">=", "<", "<=", "contains", "not_contains", "startswith", "endswith", "in", "not_in",
        "isnull", "notnull", "regex", "between"}
_AGGS = {"sum", "mean", "median", "min", "max", "count", "nunique", "std", "first", "last"}


def _col(df, name: str) -> str:
    if name in df.columns:
        return name
    low = {str(c).lower().strip(): c for c in df.columns}
    if name.lower().strip() in low:
        return low[name.lower().strip()]
    # after aggregate={'Units': 'sum'} the column is 'Units_sum': sorting by 'Units' means that
    agg = [c for c in df.columns if str(c).lower().startswith(name.lower().strip() + "_")
           and str(c).rsplit("_", 1)[-1] in _AGGS]
    if len(agg) == 1:
        return agg[0]
    raise DataError(f"no column {name!r}; columns are: {', '.join(map(str, df.columns[:40]))}")


def _coerce(series, value):
    pd = _pd()
    if pd.api.types.is_numeric_dtype(series) and isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return value
    if pd.api.types.is_datetime64_any_dtype(series):
        return pd.to_datetime(value)
    return value


def apply_filters(df, filters: Iterable[dict] | None):
    if not filters:
        return df
    pd = _pd()
    mask = pd.Series(True, index=df.index)
    for f in filters:
        if not isinstance(f, dict):
            raise DataError("each filter is {column, op, value}")
        c = _col(df, str(f.get("column", "")))
        op = str(f.get("op", "==")).lower()
        v = f.get("value")
        s = df[c]
        if op not in _OPS:
            raise DataError(f"unknown op {op!r}; use one of {', '.join(sorted(_OPS))}")
        if op in ("contains", "not_contains", "startswith", "endswith", "regex"):
            st = s.astype("string")
            if op == "contains":
                m = st.str.contains(str(v), case=bool(f.get("case", False)), regex=False)
            elif op == "not_contains":
                m = ~st.str.contains(str(v), case=bool(f.get("case", False)), regex=False)
            elif op == "startswith":
                m = st.str.startswith(str(v))
            elif op == "endswith":
                m = st.str.endswith(str(v))
            else:
                if len(str(v)) > 200:
                    raise DataError("regex too long")
                m = st.str.contains(str(v), regex=True, case=bool(f.get("case", False)))
            m = m.fillna(False).astype(bool)
        elif op in ("in", "not_in"):
            vals = v if isinstance(v, list) else [x.strip() for x in str(v).split(",")]
            vals = [_coerce(s, x) for x in vals]
            m = s.isin(vals) if op == "in" else ~s.isin(vals)
        elif op == "isnull":
            m = s.isna()
        elif op == "notnull":
            m = s.notna()
        elif op == "between":
            lo, hi = (v if isinstance(v, list) else str(v).split(","))[:2]
            m = (s >= _coerce(s, lo)) & (s <= _coerce(s, hi))
        else:
            v = _coerce(s, v)
            try:
                m = {"==": s == v, "!=": s != v, ">": s > v, ">=": s >= v, "<": s < v, "<=": s <= v}[op]
            except TypeError as exc:
                raise DataError(f"cannot compare {c} ({s.dtype}) with {v!r}: {exc}") from None
        mask &= m.fillna(False).astype(bool) if hasattr(m, "fillna") else m
    return df[mask]


def query(df, *, columns: list[str] | None = None, filters: list[dict] | None = None, group_by: list[str] | None = None,
          aggregate: dict | None = None, sort: list | None = None, limit: int | None = None, dedupe: bool = False,
          add_columns: dict | None = None):
    if add_columns:
        df = derive(df, add_columns)
    df = apply_filters(df, filters)
    if group_by:
        gb = [_col(df, g) for g in group_by]
        agg = aggregate or {}
        if not agg:
            out = df.groupby(gb, dropna=False).size().reset_index(name="count")
        else:
            spec = {}
            for col, how in agg.items():
                hows = how if isinstance(how, list) else [how]
                for h in hows:
                    if h not in _AGGS:
                        raise DataError(f"unknown aggregate {h!r}; use {', '.join(sorted(_AGGS))}")
                c = _col(df, col) if col != "*" else gb[0]
                for h in hows:
                    spec[f"{c}_{h}" if col != "*" else "count"] = (c, h if col != "*" else "count")
            out = df.groupby(gb, dropna=False).agg(**spec).reset_index()
        df = out
    elif aggregate:
        row = {}
        for col, how in aggregate.items():
            for h in (how if isinstance(how, list) else [how]):
                if h not in _AGGS:
                    raise DataError(f"unknown aggregate {h!r}")
                if col == "*":
                    row["count"] = len(df)
                else:
                    row[f"{_col(df, col)}_{h}"] = getattr(df[_col(df, col)], h)()
        df = _pd().DataFrame([row])
    if aggregate or group_by:
        # sums of decimals carry binary noise (7786.6900000000005); 9 decimals keeps real precision
        for c in df.columns:
            if _pd().api.types.is_float_dtype(df[c]):
                df[c] = df[c].round(9)
    if columns:
        df = df[[_col(df, c) for c in columns]]
    if dedupe:
        df = df.drop_duplicates()
    if sort:
        keys, asc = [], []
        for s in sort:
            if isinstance(s, dict):
                keys.append(_col(df, s.get("column", ""))), asc.append(not s.get("desc", False))
            else:
                s = str(s)
                desc = s.startswith("-")
                keys.append(_col(df, s.lstrip("-+"))), asc.append(not desc)
        df = df.sort_values(keys, ascending=asc, kind="mergesort")
    if limit:
        df = df.head(int(limit))
    return df


def describe(df, max_cols: int = 40) -> str:
    pd = _pd()
    lines = [f"{len(df):,} rows x {len(df.columns)} columns"]
    for c in list(df.columns)[:max_cols]:
        s = df[c]
        nulls = int(s.isna().sum())
        kind = str(s.dtype)
        extra = ""
        if pd.api.types.is_numeric_dtype(s) and s.notna().any():
            extra = f" min {s.min():g} max {s.max():g} mean {s.mean():.4g}" if pd.api.types.is_float_dtype(s) or \
                pd.api.types.is_integer_dtype(s) else ""
        else:
            uniq = s.nunique(dropna=True)
            top = s.astype("string").value_counts().head(3)
            extra = f" {uniq:,} distinct" + (f"; top: {', '.join(f'{k} ({v})' for k, v in top.items())}" if uniq else "")
        lines.append(f"- {c} [{kind}] nulls {nulls}{extra}"[:300])
    if len(df.columns) > max_cols:
        lines.append(f"... {len(df.columns) - max_cols} more columns")
    return "\n".join(lines)


def to_markdown(df, max_rows: int = 30, max_width: int = 40) -> str:
    cols = [str(c) for c in df.columns]
    head = df.head(max_rows)

    def cell(v):
        s = "" if v is None or (isinstance(v, float) and v != v) else str(v)
        s = s.replace("|", "\\|").replace("\n", " ")
        return s if len(s) <= max_width else s[: max_width - 1] + "…"
    out = ["| " + " | ".join(cell(c) for c in cols) + " |", "|" + "---|" * len(cols)]
    for row in head.itertuples(index=False):
        out.append("| " + " | ".join(cell(v) for v in row) + " |")
    if len(df) > max_rows:
        out.append(f"… {len(df) - max_rows:,} more rows")
    return "\n".join(out)


_FORMULA = re.compile(r"^[=+\-@\t\r]")


def write_table(df, out: str | os.PathLike, *, sheet: str = "Data") -> dict:
    o = resolve_out(out)
    ext = o.suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        df.to_csv(o, index=False, sep="\t" if ext == ".tsv" else ",", encoding="utf-8")
    elif ext in (".xlsx", ".xlsm"):
        _write_xlsx(df, o, sheet)
    elif ext == ".json":
        o.write_text(df.to_json(orient="records", indent=1, force_ascii=False, date_format="iso"), encoding="utf-8")
    elif ext in (".jsonl", ".ndjson"):
        o.write_text(df.to_json(orient="records", lines=True, force_ascii=False, date_format="iso"), encoding="utf-8")
    elif ext in (".parquet", ".pq"):
        try:
            df.to_parquet(o, index=False)
        except ImportError:
            raise DataError("Parquet needs pyarrow - pip install pyarrow") from None
    elif ext in (".md", ".markdown"):
        o.write_text(to_markdown(df, max_rows=len(df), max_width=10_000) + "\n", encoding="utf-8")
    elif ext in (".html", ".htm"):
        o.write_text(df.to_html(index=False, escape=True, border=0), encoding="utf-8")
    elif ext in (".sqlite", ".db", ".sqlite3"):
        con = sqlite3.connect(o)
        try:
            df.to_sql(sheet or "data", con, if_exists="replace", index=False)
        finally:
            con.close()
    else:
        raise DataError(f"cannot write {ext} - use csv, tsv, xlsx, json, jsonl, parquet, md, html or sqlite")
    return {"out": sandbox.rel(o), "rows": len(df), "columns": len(df.columns), "bytes": o.stat().st_size}


def _write_xlsx(df, o: Path, sheet: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    pd = _pd()
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet or "Data")[:31]
    ws.append([str(c) for c in df.columns])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for row in df.itertuples(index=False):
        vals = []
        for v in row:
            if v is None or (isinstance(v, float) and v != v) or v is pd.NA or v is pd.NaT:
                vals.append(None)
            elif hasattr(v, "to_pydatetime"):
                vals.append(v.to_pydatetime())
            elif hasattr(v, "item") and not isinstance(v, (str, bytes)):
                vals.append(v.item())
            else:
                vals.append(v if isinstance(v, (int, float, bool)) else str(v))
        ws.append(vals)
    # untrusted text that looks like a formula stays text (CSV/formula injection)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if isinstance(c.value, str) and _FORMULA.match(c.value):
                c.data_type = "s"
    for i, col in enumerate(df.columns, 1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist()])
        ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60)
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    wb.save(o)


# -- structured files ----------------------------------------------------------------------------
def load_structured(p: Path) -> Any:
    k = kind_of(p)
    if p.stat().st_size > MAX_LOAD():
        raise DataError(f"{p.name} is too big to parse in memory")
    text = p.read_text(encoding=detect_encoding(p), errors="replace")
    if k == "json":
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            raise DataError(f"invalid JSON at line {exc.lineno} col {exc.colno}: {exc.msg}") from None
    if k == "jsonl":
        out = []
        for i, line in enumerate(text.splitlines(), 1):
            if line.strip():
                try:
                    out.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise DataError(f"invalid JSON on line {i}: {exc.msg}") from None
        return out
    if k == "yaml":
        try:
            import yaml
        except ImportError:
            raise DataError("YAML needs PyYAML - pip install pyyaml") from None
        try:
            docs = list(yaml.safe_load_all(text))  # safe_load: no Python objects, no code
        except yaml.YAMLError as exc:
            raise DataError(f"invalid YAML: {str(exc)[:300]}") from None
        return docs[0] if len(docs) == 1 else docs
    if k == "toml":
        try:
            import tomllib
        except ImportError:  # Python 3.10
            try:
                import tomli as tomllib  # type: ignore[no-redef]
            except ImportError:
                raise DataError("TOML on Python 3.10 needs tomli - pip install tomli") from None
        try:
            return tomllib.loads(text)
        except Exception as exc:
            raise DataError(f"invalid TOML: {exc}") from None
    if k == "xml":
        return xml_to_obj(text)
    raise DataError(f"{p.name}: not JSON/JSONL/YAML/TOML/XML")


def xml_to_obj(text: str) -> dict:
    from defusedxml import ElementTree as ET  # blocks entity expansion (billion laughs), external entities and DTDs

    try:
        root = ET.fromstring(text, forbid_dtd=True)
    except ET.ParseError as exc:
        raise DataError(f"invalid XML: {exc}") from None
    except Exception as exc:  # defusedxml raises DTDForbidden / EntitiesForbidden / ExternalReferenceForbidden
        raise DataError(f"unsafe XML refused ({exc.__class__.__name__}): DTDs and entities are not allowed") from None

    def strip(tag: str) -> str:
        return tag.split("}", 1)[1] if tag.startswith("{") else tag

    def conv(el, depth=0):
        if depth > 60:
            return "…"
        d: dict = {f"@{strip(k)}": v for k, v in el.attrib.items()}
        kids = list(el)
        for k in kids:
            v = conv(k, depth + 1)
            t = strip(k.tag)
            if t in d:
                if not isinstance(d[t], list):
                    d[t] = [d[t]]
                d[t].append(v)
            else:
                d[t] = v
        text = (el.text or "").strip()
        if text:
            if not d:
                return text
            d["#text"] = text
        return d
    return {strip(root.tag): conv(root)}


def outline(obj: Any, depth: int = 0, max_depth: int = 3, max_items: int = 12) -> list[str]:
    pad = "  " * depth
    out = []
    if isinstance(obj, dict):
        for i, (k, v) in enumerate(obj.items()):
            if i >= max_items:
                out.append(f"{pad}… {len(obj) - max_items} more keys")
                break
            t = type(v).__name__
            n = f" ({len(v)})" if isinstance(v, (list, dict)) else f" = {str(v)[:60]!r}" if not isinstance(v, (dict, list)) else ""
            out.append(f"{pad}{k}: {t}{n}")
            if isinstance(v, (dict, list)) and depth + 1 < max_depth:
                out += outline(v, depth + 1, max_depth, max_items)
    elif isinstance(obj, list):
        if obj:
            out.append(f"{pad}[{len(obj)} items, first is {type(obj[0]).__name__}]")
            if isinstance(obj[0], (dict, list)) and depth + 1 < max_depth:
                out += outline(obj[0], depth + 1, max_depth, max_items)
    return out


def select(obj: Any, path: str) -> Any:
    """Tiny dotted path: a.b[0].c / items.*.name"""
    if not path:
        return obj
    parts = re.findall(r"[^.\[\]]+|\[\d+\]", path)
    cur = [obj]
    for part in parts:
        nxt = []
        for o in cur:
            if part.startswith("["):
                i = int(part[1:-1])
                if isinstance(o, list) and -len(o) <= i < len(o):
                    nxt.append(o[i])
            elif part == "*":
                nxt += list(o.values()) if isinstance(o, dict) else list(o) if isinstance(o, list) else []
            elif isinstance(o, dict) and part in o:
                nxt.append(o[part])
            elif isinstance(o, list) and part.isdigit() and int(part) < len(o):
                nxt.append(o[int(part)])
        cur = nxt
    return cur[0] if len(cur) == 1 and "*" not in parts else cur


def write_structured(obj: Any, out: str | os.PathLike) -> dict:
    o = resolve_out(out)
    ext = o.suffix.lower()
    if ext == ".json":
        o.write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    elif ext in (".jsonl", ".ndjson"):
        items = obj if isinstance(obj, list) else [obj]
        o.write_text("".join(json.dumps(x, ensure_ascii=False, default=str) + "\n" for x in items), encoding="utf-8")
    elif ext in (".yaml", ".yml"):
        import yaml

        o.write_text(yaml.safe_dump(obj, sort_keys=False, allow_unicode=True), encoding="utf-8")
    elif ext == ".toml":
        try:
            import tomli_w
        except ImportError:
            o.write_text(_toml_dump(obj), encoding="utf-8")
        else:
            o.write_text(tomli_w.dumps(obj), encoding="utf-8")
    else:
        raise DataError(f"cannot write structured data as {ext} (json, jsonl, yaml, toml)")
    return {"out": sandbox.rel(o), "bytes": o.stat().st_size}


def _toml_dump(obj: Any) -> str:
    """Minimal TOML writer (tables, arrays of scalars/tables, strings, numbers, booleans)."""
    if not isinstance(obj, dict):
        raise DataError("TOML needs a mapping at the top level")

    def val(v):
        if isinstance(v, bool):
            return "true" if v else "false"
        if isinstance(v, (int, float)):
            return repr(v)
        if isinstance(v, list) and all(not isinstance(x, dict) for x in v):
            return "[" + ", ".join(val(x) for x in v) + "]"
        if v is None:
            return '""'
        return json.dumps(str(v), ensure_ascii=False)

    lines: list[str] = []

    def emit(d: dict, prefix: str):
        simple = {k: v for k, v in d.items() if not isinstance(v, dict) and not (isinstance(v, list) and v and all(isinstance(x, dict) for x in v))}
        for k, v in simple.items():
            lines.append(f"{_tk(k)} = {val(v)}")
        for k, v in d.items():
            if isinstance(v, dict):
                lines.append(f"\n[{prefix}{_tk(k)}]")
                emit(v, f"{prefix}{_tk(k)}.")
            elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                for x in v:
                    lines.append(f"\n[[{prefix}{_tk(k)}]]")
                    emit(x, f"{prefix}{_tk(k)}.")
    emit(obj, "")
    return "\n".join(lines).lstrip("\n") + "\n"


def _tk(k: str) -> str:
    return k if re.match(r"^[A-Za-z0-9_-]+$", str(k)) else json.dumps(str(k))


# -- SQLite (read-only) ------------------------------------------------------------------------------
def sqlite_connect(p: Path) -> sqlite3.Connection:
    uri = p.resolve().as_uri() + "?mode=ro&immutable=1"  # file:///C:/… on Windows, percent-encoded
    con = sqlite3.connect(uri, uri=True, timeout=5, check_same_thread=False)
    allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION}
    pragma_ok = {"table_info", "table_list", "index_list", "index_info", "foreign_key_list", "table_xinfo"}

    def auth(action, arg1, arg2, dbname, source):
        if action in allowed:
            return sqlite3.SQLITE_OK
        if action == sqlite3.SQLITE_PRAGMA and str(arg1).lower() in pragma_ok:
            return sqlite3.SQLITE_OK
        return sqlite3.SQLITE_DENY
    con.set_authorizer(auth)
    deadline = time.monotonic() + 30
    con.set_progress_handler(lambda: 1 if time.monotonic() > deadline else 0, 10000)
    return con


def sqlite_schema(p: Path) -> list[dict]:
    con = sqlite_connect(p)
    try:
        out = []
        for name, typ in con.execute("SELECT name, type FROM sqlite_master WHERE type IN ('table','view') "
                                     "AND name NOT LIKE 'sqlite_%' ORDER BY name"):
            cols = [f"{r[1]} {r[2]}".strip() for r in con.execute(f'PRAGMA table_info("{name}")')]
            try:
                n = con.execute(f'SELECT COUNT(*) FROM "{name}"').fetchone()[0]  # nosec B608 - name from sqlite_master
            except sqlite3.DatabaseError:
                n = None
            out.append({"name": name, "type": typ, "rows": n, "columns": cols})
        return out
    finally:
        con.close()


def sqlite_query(path: str | os.PathLike, sql: str, *, params: list | None = None, limit: int = 200) -> tuple[list[str], list[tuple], bool]:
    p = resolve_in(path)
    if kind_of(p) != "sqlite":
        raise DataError(f"{p.name} is not an SQLite database")
    s = sql.strip().rstrip(";")
    if ";" in s:
        raise DataError("one statement at a time")
    if not re.match(r"(?is)^\s*(select|with|pragma\s+(table_info|table_list|index_list|foreign_key_list))\b", s):
        raise DataError("read-only: only SELECT / WITH … SELECT / PRAGMA table_info are allowed")
    con = sqlite_connect(p)
    try:
        try:
            cur = con.execute(s, params or [])
        except sqlite3.DatabaseError as exc:
            raise DataError(f"SQLite: {exc}") from None
        cols = [d[0] for d in cur.description or []]
        rows = cur.fetchmany(int(limit) + 1)
        return cols, rows[:limit], len(rows) > limit
    finally:
        con.close()


# -- text / logs -------------------------------------------------------------------------------------
def tail(path: str | os.PathLike, lines: int = 50) -> str:
    p = resolve_in(path)
    n = max(1, min(int(lines or 50), 2000))
    enc = detect_encoding(p)
    with p.open("rb") as fh:
        fh.seek(0, 2)
        size = fh.tell()
        block, data = 64 * 1024, b""
        pos = size
        while pos > 0 and data.count(b"\n") <= n:
            step = min(block, pos)
            pos -= step
            fh.seek(pos)
            data = fh.read(step) + data
            if len(data) > 8 * 1024 * 1024:
                break
    out = data.decode(enc, errors="replace").splitlines()[-n:]
    return "\n".join(ln[:2000] for ln in out)


def grep(path: str | os.PathLike, pattern: str, *, regex: bool = False, ignore_case: bool = True, context: int = 0,
         max_matches: int = 200, invert: bool = False) -> tuple[list[str], int, int]:
    p = resolve_in(path)
    if regex and len(pattern) > 300:
        raise DataError("pattern too long")
    flags = re.I if ignore_case else 0
    rx = re.compile(pattern if regex else re.escape(pattern), flags)
    ctx = max(0, min(int(context or 0), 10))
    before: deque = deque(maxlen=ctx)
    out, total, after, lineno = [], 0, 0, 0
    with p.open("r", encoding=detect_encoding(p), errors="replace") as fh:
        for lineno, line in enumerate(fh, 1):
            line = line.rstrip("\n")[:4000]
            hit = bool(rx.search(line)) != invert
            if hit:
                total += 1
                if len(out) < max_matches:
                    out += [f"{lineno - len(before) + i}- {b}" for i, b in enumerate(before)]
                    out.append(f"{lineno}: {line}")
                    after = ctx
                before.clear()
            elif after and len(out) < max_matches + ctx * 2:
                out.append(f"{lineno}- {line}")
                after -= 1
            else:
                before.append(line)
    return out, total, lineno


LOG_LEVEL = re.compile(r"\b(TRACE|DEBUG|INFO|NOTICE|WARN(?:ING)?|ERROR|ERR|CRIT(?:ICAL)?|FATAL|ALERT|EMERG)\b")


def log_summary(path: str | os.PathLike) -> dict:
    p = resolve_in(path)
    levels: dict[str, int] = {}
    first = last = None
    n = 0
    ts = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2}")
    with p.open("r", encoding=detect_encoding(p), errors="replace") as fh:
        for n, line in enumerate(fh, 1):
            m = LOG_LEVEL.search(line[:300])
            if m:
                lv = m.group(1).upper().replace("WARNING", "WARN").replace("ERR", "ERROR").replace("ERROROR", "ERROR")
                levels[lv] = levels.get(lv, 0) + 1
            t = ts.search(line[:100])
            if t:
                first = first or t.group(0)
                last = t.group(0)
    return {"lines": n, "levels": levels, "first_timestamp": first, "last_timestamp": last}


# -- images / PDF ------------------------------------------------------------------------------------
def image_info(path: str | os.PathLike) -> dict:
    from PIL import ExifTags, Image

    p = resolve_in(path)
    Image.MAX_IMAGE_PIXELS = 120_000_000  # decompression-bomb guard (raises above 2x this)
    try:
        with Image.open(p) as im:
            info = {"format": im.format, "size": list(im.size), "mode": im.mode, "frames": getattr(im, "n_frames", 1),
                    "bytes": p.stat().st_size}
            dpi = im.info.get("dpi")
            if dpi:
                info["dpi"] = [round(float(x)) for x in dpi]
            if im.info.get("icc_profile"):
                info["icc_profile"] = True
            exif = im.getexif()
            if exif:
                tags = {}
                for k, v in exif.items():
                    name = ExifTags.TAGS.get(k, str(k))
                    if name in ("Make", "Model", "DateTime", "Software", "Orientation", "Artist", "Copyright",
                                "ImageDescription"):
                        tags[name] = str(v)[:120]
                gps = exif.get_ifd(0x8825) if hasattr(exif, "get_ifd") else {}
                if gps:
                    tags["GPS"] = "present - the photo contains a location; strip EXIF before sharing"
                info["exif"] = tags
    except Image.DecompressionBombError as exc:
        raise DataError(f"image refused: {exc}") from None
    except OSError as exc:
        raise DataError(f"not a readable image: {exc}") from None
    return info


def pdf_text(path: str | os.PathLike, first: int = 1, last: int = 0, layout: bool = False, max_chars: int = 40000) -> dict:
    p = resolve_in(path)
    tool = shutil.which("pdftotext")
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(p))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise DataError("the PDF is password-protected")
        pages = len(reader.pages)
        meta = {k.lstrip("/"): str(v)[:200] for k, v in (reader.metadata or {}).items()}
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"not a readable PDF: {exc.__class__.__name__}: {exc}") from None
    first = max(1, int(first or 1))
    last = min(pages, int(last or pages))
    if tool:
        args = [tool, "-f", str(first), "-l", str(last), "-enc", "UTF-8"] + (["-layout"] if layout else []) + [str(p), "-"]
        r = subprocess.run(args, capture_output=True, timeout=120, env=sandbox.clean_env())
        text = r.stdout.decode("utf-8", errors="replace")
        how = "pdftotext"
    else:
        text = "\n\f".join((reader.pages[i].extract_text() or "") for i in range(first - 1, last))
        how = "pypdf"
    chars = len(text)
    if chars > max_chars:
        text = text[:max_chars] + f"\n[… {chars - max_chars} more chars - ask for later pages]"
    if not text.strip():
        text = "(no text layer - this looks like a scanned PDF; OCR is needed)"
    return {"pages": pages, "range": [first, last], "engine": how, "metadata": meta, "text": text}


def markdown_outline(text: str) -> list[str]:
    out, fence = [], False
    for line in text.splitlines():
        if line.strip().startswith(("```", "~~~")):
            fence = not fence
        elif not fence:
            m = re.match(r"^(#{1,6})\s+(.*)", line)
            if m:
                out.append("  " * (len(m.group(1)) - 1) + m.group(2).strip())
    return out
