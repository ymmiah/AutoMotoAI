"""Builders for Word (.docx), Excel (.xlsx) and PowerPoint (.pptx) from a JSON spec.

The agent describes the document as data (blocks, sheets, slides); these functions lay it out
with sensible publication defaults (A4, real heading styles, captions, a working table of contents,
page numbers, Harvard references, number formats, charts, speaker notes).
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Any

from automoto import sandbox
from automoto.office import OfficeError, inspect_file, lo_available, lo_convert

IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
MAX_IMG = 25 * 1024 * 1024


def _hex(c: str | None, default: str) -> str:
    c = (c or default).lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", c):
        raise OfficeError(f"colour must be like #1F4E79, got {c!r}")
    return c.upper()


def _image(path: str) -> Path:
    p = sandbox.resolve(path, allow_outside=True)
    if not p.is_file():
        raise OfficeError(f"image not found: {path}")
    if p.suffix.lower() not in IMG_EXT:
        raise OfficeError(f"images must be PNG/JPEG/GIF/BMP/TIFF (convert SVG to PNG first): {p.name}")
    if p.stat().st_size > MAX_IMG:
        raise OfficeError(f"image too large: {p.name}")
    return p


# ======================================================================================= Word
def _field(paragraph, instr: str, placeholder: str = "1"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r

    p = paragraph._p
    p.append(fld("begin"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    r.append(t)
    p.append(r)
    p.append(fld("separate"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = placeholder
    r.append(t)
    p.append(r)
    p.append(fld("end"))


def _hyperlink(paragraph, url: str, text: str, color: str = "0563C1"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    if not re.match(r"^(https?://|mailto:)", url):
        paragraph.add_run(text)
        return
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _inline(paragraph, text: str, *, size=None):
    """Add text with **bold**, *italic*, `code` and [link](url) markup."""
    runs = []
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            _hyperlink(paragraph, m.group(2), m.group(1))
            continue
        else:
            r = paragraph.add_run(part)
        if size:
            r.font.size = size
        runs.append(r)
    return runs


def harvard(ref: dict[str, Any] | str) -> str:
    """Format one reference in Harvard (Cite Them Right) style, with *italic* markup for the title."""
    if isinstance(ref, str):
        return ref
    authors = ref.get("authors") or ref.get("author") or ref.get("organisation") or "Anon."
    if isinstance(authors, list):
        authors = authors[0] if len(authors) == 1 else ", ".join(authors[:-1]) + " and " + authors[-1]
    year = ref.get("year", "n.d.")
    title = ref.get("title", "Untitled")
    kind = (ref.get("type") or ("journal" if ref.get("journal") else "web" if ref.get("url") and not ref.get("publisher") else "book")).lower()
    if kind == "journal":
        s = f"{authors} ({year}) '{title}', *{ref.get('journal', '')}*"
        if ref.get("volume"):
            s += f", {ref['volume']}"
            if ref.get("issue"):
                s += f"({ref['issue']})"
        if ref.get("pages"):
            s += f", pp. {ref['pages']}"
        s += "."
        if ref.get("doi"):
            s += f" doi: {ref['doi']}."
    elif kind == "web":
        s = f"{authors} ({year}) *{title}*."
    else:
        s = f"{authors} ({year}) *{title}*."
        if ref.get("edition"):
            s += f" {ref['edition']} edn."
        if ref.get("place") and ref.get("publisher"):
            s += f" {ref['place']}: {ref['publisher']}."
        elif ref.get("publisher"):
            s += f" {ref['publisher']}."
    if ref.get("url"):
        s += f" Available at: {ref['url']}"
        if ref.get("accessed"):
            s += f" (Accessed: {ref['accessed']})"
        s += "."
    return s


_LABEL = re.compile(r"^\s*(figure|fig\.|table|source)\s*\d*\s*[:.\-–]\s*", re.I)


def _strip_label(text: Any) -> str:
    """'Figure 1: Wind output' -> 'Wind output' (the builder numbers captions itself)."""
    return _LABEL.sub("", str(text or ""), count=1).strip()


def _sort_key(r: Any) -> str:
    return (harvard(r) if not isinstance(r, str) else r).lower()


def build_docx(spec: dict[str, Any], out: Path, *, toc_pages: dict[str, int] | None = None) -> dict[str, Any]:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    if not isinstance(spec, dict):
        raise OfficeError("spec must be an object with 'blocks'")
    blocks = spec.get("blocks") or []
    if not isinstance(blocks, list) or len(blocks) > 3000:
        raise OfficeError("blocks must be a list (max 3000)")
    d = docx.Document()
    theme = spec.get("theme") or {}
    font = theme.get("font", "Calibri")
    head_font = theme.get("heading_font", font)
    size = float(theme.get("size", 11))
    color = RGBColor.from_string(_hex(theme.get("heading_color"), "1F3864"))
    st = d.styles
    st["Normal"].font.name = font
    st["Normal"].font.size = Pt(size)
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), font)
    pf = st["Normal"].paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = float(theme.get("line_spacing", 1.15))
    for lvl, sz in ((1, size + 7), (2, size + 3), (3, size + 1)):
        h = st[f"Heading {lvl}"]
        h.font.name, h.font.size, h.font.color.rgb, h.font.bold = head_font, Pt(sz), color, True
        rpr = h.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(k), head_font)
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(k)) is not None:
                del rfonts.attrib[qn(k)]
        h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    for name in ("Title", "Subtitle", "Caption", "TOC Heading"):
        try:
            st[name].font.name = head_font
        except KeyError:
            pass
    st["Caption"].font.size = Pt(size - 1)
    st["Caption"].font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    sec = d.sections[0]
    page = (spec.get("page") or {})
    if str(page.get("size", "A4")).lower() == "letter":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if page.get("orientation") == "landscape":
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    m = float(page.get("margins_cm", 2.54))
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(m)
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    meta = spec.get("properties") or {}
    cp = d.core_properties
    cp.title = str(meta.get("title") or spec.get("title") or "")[:250]
    cp.author = str(meta.get("author") or "")[:250]
    cp.subject = str(meta.get("subject") or "")[:250]
    cp.keywords = str(meta.get("keywords") or "")[:250]
    lang = meta.get("language", "en-GB")
    st["Normal"].element.get_or_add_rPr()
    lg = OxmlElement("w:lang")
    lg.set(qn("w:val"), lang)
    st["Normal"].element.rPr.append(lg)

    has_title_page = any(isinstance(b, dict) and b.get("type") == "title_page" for b in blocks)
    if has_title_page:
        sec.different_first_page_header_footer = True
    if spec.get("header"):
        hp = sec.header.paragraphs[0]
        hp.text = str(spec["header"])
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(size - 2)
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if spec.get("page_numbers", True):
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if spec.get("footer"):
            fp.add_run(str(spec["footer"]) + "   ·   ")
        style = spec.get("page_number_style", "Page {PAGE} of {NUMPAGES}")
        for part in re.split(r"(\{PAGE\}|\{NUMPAGES\})", style):
            if part == "{PAGE}":
                _field(fp, "PAGE", "1")
            elif part == "{NUMPAGES}":
                _field(fp, "NUMPAGES", "1")
            elif part:
                fp.add_run(part)
    elif spec.get("footer"):
        sec.footer.paragraphs[0].text = str(spec["footer"])

    counters = {"fig": 0, "tab": 0}
    headings: list[tuple[int, str]] = []
    toc_levels = 3

    def tab_stop(p):
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    def _refs_heading_before(k: int) -> bool:
        """The author already wrote a 'References' heading right before the references block."""
        prev = blocks[k - 1] if k > 0 else None
        return isinstance(prev, dict) and prev.get("type") == "heading" and bool(
            re.search(r"\b(references|bibliography|works cited|reference list)\b", str(prev.get("text", "")), re.I))

    # pre-scan headings for the TOC (so the first pass already has the right number of lines)
    for k, b in enumerate(blocks):
        if isinstance(b, dict) and b.get("type") == "heading" and int(b.get("level", 1)) <= toc_levels:
            headings.append((int(b.get("level", 1)), str(b.get("text", ""))))
        elif isinstance(b, dict) and b.get("type") == "references" and not _refs_heading_before(k):
            headings.append((int(b.get("level", 1)), str(b.get("title", "References"))))

    for k, b in enumerate(blocks):
        if not isinstance(b, dict):
            b = {"type": "paragraph", "text": str(b)}
        t = b.get("type", "paragraph")
        if t == "title_page":
            for _ in range(6):
                d.add_paragraph()
            p = d.add_paragraph(style="Title")
            p.add_run(str(b.get("title", cp.title or "")))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if b.get("subtitle"):
                p = d.add_paragraph(str(b["subtitle"]), style="Subtitle")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for _ in range(4):
                d.add_paragraph()
            for key in ("author", "student_id", "course", "module", "institution", "supervisor", "date", "word_count"):
                if b.get(key):
                    label = {"student_id": "Student ID", "word_count": "Word count"}.get(key, key.replace("_", " ").title())
                    p = d.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if key in ("author", "institution", "date"):
                        p.add_run(str(b[key])).font.size = Pt(size + 2)
                    else:
                        p.add_run(f"{label}: {b[key]}")
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "heading":
            lvl = max(1, min(int(b.get("level", 1)), 6))
            d.add_heading(str(b.get("text", "")), level=lvl)
        elif t == "paragraph":
            p = d.add_paragraph(style=b.get("style")) if b.get("style") in [s.name for s in st] else d.add_paragraph()
            runs = _inline(p, b.get("text", ""))
            if b.get("align") in ("center", "right", "justify"):
                p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                               "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[b["align"]]
            elif theme.get("justify"):
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if b.get("comment") and runs and hasattr(d, "add_comment"):
                d.add_comment(runs, text=str(b["comment"]), author=str(b.get("comment_author", "AutoMoto")), initials="AM")
        elif t in ("bullets", "numbered"):
            base = "List Bullet" if t == "bullets" else "List Number"
            for item in b.get("items", []):
                lvl = 0
                if isinstance(item, (list, tuple)):
                    item, lvl = item[0], int(item[1])
                name = base if lvl == 0 else f"{base} {min(lvl + 1, 3)}"
                p = d.add_paragraph(style=name)
                _inline(p, item)
        elif t == "quote":
            p = d.add_paragraph(style="Quote") if "Quote" in [s.name for s in st] else d.add_paragraph()
            _inline(p, b.get("text", ""))
        elif t == "table":
            rows = b.get("rows") or []
            if not rows or not all(isinstance(r, list) for r in rows):
                raise OfficeError("table needs rows: [[...], [...]]")
            ncols = max(len(r) for r in rows)
            if b.get("caption"):
                counters["tab"] += 1
                cap = d.add_paragraph(style="Caption")
                cap.add_run(f"Table {counters['tab']}: ").bold = True
                cap.add_run(_strip_label(b["caption"]))
                cap.paragraph_format.keep_with_next = True
            tbl = d.add_table(rows=len(rows), cols=ncols)
            sty = b.get("style", "Light Grid Accent 1")
            try:
                tbl.style = st[sty]
            except KeyError:
                tbl.style = st["Table Grid"]
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            look = tbl._tbl.tblPr.find(qn("w:tblLook"))
            if look is None:
                look = OxmlElement("w:tblLook")
                tbl._tbl.tblPr.append(look)
            for k, v in (("w:firstRow", "1"), ("w:firstColumn", "0"), ("w:lastRow", "0"), ("w:lastColumn", "0"),
                         ("w:noHBand", "0"), ("w:noVBand", "1"), ("w:val", "0420")):
                look.set(qn(k), v)
            for i, r in enumerate(rows):
                for j in range(ncols):
                    v = r[j] if j < len(r) else ""
                    cell = tbl.cell(i, j)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    _inline(para, "" if v is None else str(v), size=Pt(size - 1))
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.space_before = Pt(2)
                    if i > 0 and j > 0 and re.fullmatch(r"[-+£$€]?[\d,]+(\.\d+)?%?", str(v).strip()):
                        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if i == 0 and b.get("header", True):
                        for run in para.runs:
                            run.bold = True
            if b.get("header", True):
                trpr = tbl.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                th.set(qn("w:val"), "true")
                trpr.append(th)
            if b.get("source"):
                p = d.add_paragraph(style="Caption")
                p.add_run(f"Source: {_strip_label(b['source'])}")
            else:
                d.add_paragraph()
        elif t == "image":
            img = _image(str(b.get("path", "")))
            width = Cm(float(b.get("width_cm", 15)))
            if width > text_width:
                width = text_width
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = bool(b.get("caption"))
            pic = p.add_run().add_picture(str(img), width=width)
            alt = str(b.get("alt") or b.get("caption") or img.stem)
            pic._inline.docPr.set("descr", alt[:500])
            pic._inline.docPr.set("title", alt[:120])
            if b.get("caption"):
                counters["fig"] += 1
                cap = d.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run(f"Figure {counters['fig']}: ").bold = True
                cap.add_run(_strip_label(b["caption"]))
        elif t == "toc":
            toc_levels = int(b.get("levels", 3))
            p = d.add_paragraph(str(b.get("title", "Contents")), style="TOC Heading")
            entries = [(lv, tx) for lv, tx in headings if lv <= toc_levels]
            first = None
            for i, (lv, tx) in enumerate(entries):
                ep = d.add_paragraph()
                try:
                    ep.style = st[f"toc {lv}"]
                except KeyError:
                    ep.paragraph_format.left_indent = Cm(0.6 * (lv - 1))
                ep.paragraph_format.space_after = Pt(2)
                tab_stop(ep)
                if i == 0:
                    first = ep
                    from docx.oxml import OxmlElement as OE

                    def _fc(kind):
                        r = OE("w:r")
                        f = OE("w:fldChar")
                        f.set(qn("w:fldCharType"), kind)
                        r.append(f)
                        return r
                    ep._p.append(_fc("begin"))
                    r = OE("w:r")
                    it = OE("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = f' TOC \\o "1-{toc_levels}" \\h \\z \\u '
                    r.append(it)
                    ep._p.append(r)
                    ep._p.append(_fc("separate"))
                run = ep.add_run(tx)
                if lv == 1:
                    run.bold = True
                pg = (toc_pages or {}).get(tx)
                ep.add_run("\t" + (str(pg) if pg else "0"))
                if i == len(entries) - 1:
                    from docx.oxml import OxmlElement as OE2

                    r = OE2("w:r")
                    f = OE2("w:fldChar")
                    f.set(qn("w:fldCharType"), "end")
                    r.append(f)
                    ep._p.append(r)
            if first is None:
                d.add_paragraph("(no headings yet)")
            if b.get("page_break", True):
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "page_break":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "references":
            if not _refs_heading_before(k):
                d.add_heading(str(b.get("title", "References")), level=int(b.get("level", 1)))
            items = list(b.get("items") or [])
            if b.get("sort", True):
                items.sort(key=_sort_key)
            for ref in items:
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)
                p.paragraph_format.space_after = Pt(6)
                _inline(p, harvard(ref))
        else:
            raise OfficeError(f"unknown block type '{t}'. Use title_page, heading, paragraph, bullets, numbered, "
                              "quote, table, image, toc, page_break, references")
    d.save(str(out))
    return {"headings": headings, "figures": counters["fig"], "tables": counters["tab"],
            "has_toc": any(isinstance(b, dict) and b.get("type") == "toc" for b in blocks)}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def toc_page_numbers(pdf: Path, headings: list[tuple[int, str]]) -> dict[str, int]:
    """Find the page each heading starts on in a rendered PDF.

    Headings are matched as whole lines (a heading word that also appears in body text on an
    earlier page must not win), searching forward from the previous heading's page.
    """
    from pypdf import PdfReader

    raw = [(pg.extract_text() or "") for pg in PdfReader(str(pdf)).pages]
    lines = [[_norm(ln) for ln in t.splitlines() if ln.strip()] for t in raw]
    texts = [_norm(t) for t in raw]
    # skip the contents page(s): the pages that list most headings next to "contents"
    start = 0
    for i, txt in enumerate(texts[:6]):
        if "contents" in txt and sum(1 for _, h in headings if _norm(h) in txt) >= max(2, len(headings) // 2):
            start = i + 1
    found: dict[str, int] = {}
    cur = start
    for _, h in headings:
        key = _norm(h)
        if not key:
            continue
        hit = next((i for i in range(cur, len(lines)) if any(ln == key or (ln.startswith(key) and len(ln) <= len(key) + 4)
                                                             for ln in lines[i])), None)
        if hit is None:  # long headings may wrap across lines: fall back to the joined text
            hit = next((i for i in range(cur, len(texts)) if key in texts[i]), None)
        if hit is not None:
            found[h] = hit + 1
            cur = hit
    return found


def create_docx(spec: dict[str, Any], out: Path) -> str:
    info = build_docx(spec, out)
    notes = []
    if info["has_toc"] and lo_available():
        with tempfile.TemporaryDirectory(prefix="automoto-toc-") as tmp:
            pdf = lo_convert(out, "pdf", Path(tmp))
            pages = toc_page_numbers(pdf, info["headings"])
        build_docx(spec, out, toc_pages=pages)
        missing = [h for _, h in info["headings"] if h not in pages]
        notes.append(f"table of contents filled with page numbers ({len(pages)}/{len(info['headings'])} headings found)"
                     + (f"; not found: {', '.join(missing[:5])}" if missing else ""))
    elif info["has_toc"]:
        notes.append("TOC field inserted - Word fills page numbers when you press F9 (LibreOffice not found)")
    return (f"wrote {sandbox.rel(out)}: {len(info['headings'])} headings, {info['tables']} table(s), "
            f"{info['figures']} figure(s)" + ("; " + "; ".join(notes) if notes else ""))


def markdown_to_spec(md: str) -> dict[str, Any]:
    """Small Markdown → docx spec converter (headings, paragraphs, lists, tables, images, quotes)."""
    blocks: list[dict[str, Any]] = []
    lines = md.splitlines()
    i = 0
    para: list[str] = []

    def flush():
        if para:
            blocks.append({"type": "paragraph", "text": " ".join(x.strip() for x in para)})
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush()
        elif m := re.match(r"^(#{1,6})\s+(.*)", s):
            flush()
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
        elif re.match(r"^[-*+]\s+", s) or re.match(r"^\d+[.)]\s+", s):
            flush()
            kind = "bullets" if re.match(r"^[-*+]\s+", s) else "numbered"
            items = []
            same = r"^[-*+]\s+" if kind == "bullets" else r"^\d+[.)]\s+"
            while i < len(lines) and re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]) and (
                    lines[i][:1].isspace() or re.match(same, lines[i])):
                raw = lines[i]
                lvl = (len(raw) - len(raw.lstrip())) // 2
                items.append([re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", raw), min(lvl, 2)])
                i += 1
            blocks.append({"type": kind, "items": items})
            continue
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue
        elif m := re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s):
            flush()
            blocks.append({"type": "image", "path": m.group(2), "caption": m.group(1) or None})
        elif s.startswith(">"):
            flush()
            blocks.append({"type": "quote", "text": s.lstrip("> ")})
        elif s in ("---", "***", "\\pagebreak"):
            flush()
            blocks.append({"type": "page_break"})
        else:
            para.append(s)
        i += 1
    flush()
    return {"blocks": blocks}


# ======================================================================================= Excel
_RANGE = re.compile(r"^([A-Z]{1,3})(\d+)(:([A-Z]{1,3})(\d+))?$")


def create_xlsx(spec: dict[str, Any], out: Path) -> str:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sheets = spec.get("sheets") if isinstance(spec, dict) else None
    if not sheets or not isinstance(sheets, list) or len(sheets) > 50:
        raise OfficeError("spec needs sheets: [{name, rows, ...}] (max 50)")
    theme = spec.get("theme") or {}
    head_fill = _hex(theme.get("header_fill"), "1F4E79")
    font_name = theme.get("font", "Calibri")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    report = []
    for sh in sheets:
        name = str(sh.get("name", f"Sheet{len(wb.sheetnames) + 1}"))[:31]
        if re.search(r"[\[\]:*?/\\]", name):
            raise OfficeError(f"sheet name {name!r} has a character Excel doesn't allow")
        ws = wb.create_sheet(name)
        rows = sh.get("rows") or []
        if len(rows) > 100000:
            raise OfficeError("max 100000 rows per sheet")
        start = sh.get("start", "A1")
        m = _RANGE.match(start)
        r0 = int(m.group(2)) if m else 1
        c0 = openpyxl.utils.column_index_from_string(m.group(1)) if m else 1
        nform = 0
        for i, row in enumerate(rows):
            if not isinstance(row, list):
                raise OfficeError("each row must be a list")
            for j, v in enumerate(row):
                cell = ws.cell(row=r0 + i, column=c0 + j, value=v)
                cell.font = Font(name=font_name, size=11)
                if isinstance(v, str) and v.startswith("="):
                    nform += 1
        header = sh.get("header", True) and rows
        ncols = max((len(r) for r in rows), default=0)
        if header:
            for j in range(ncols):
                c = ws.cell(row=r0, column=c0 + j)
                c.font = Font(name=font_name, bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=head_fill)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if sh.get("borders", True) and rows:
            for i in range(len(rows)):
                for j in range(ncols):
                    ws.cell(row=r0 + i, column=c0 + j).border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for rng, fmt in (sh.get("number_formats") or {}).items():
            for row in ws[rng] if ":" in rng else [[ws[rng]]]:
                for c in row:
                    c.number_format = fmt
        for rng in sh.get("bold") or []:
            for row in ws[rng] if ":" in rng else [[ws[rng]]]:
                for c in row:
                    c.font = Font(name=font_name, bold=True)
        widths = sh.get("col_widths")
        if widths == "auto" or widths is None:
            for j in range(ncols):
                letter = get_column_letter(c0 + j)
                longest = max((len(str(r[j])) for r in rows if j < len(r) and r[j] is not None
                               and not (isinstance(r[j], str) and r[j].startswith("="))), default=8)
                ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)
        elif isinstance(widths, dict):
            for letter, w in widths.items():
                ws.column_dimensions[letter].width = float(w)
        if sh.get("freeze"):
            ws.freeze_panes = sh["freeze"]
        elif header:
            ws.freeze_panes = ws.cell(row=r0 + 1, column=c0)
        if sh.get("autofilter") and rows:
            ws.auto_filter.ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(c0 + ncols - 1)}{r0 + len(rows) - 1}"
        for cf in sh.get("conditional") or []:
            kind, rng = cf.get("type"), cf.get("range")
            if kind == "color_scale":
                ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=_hex(cf.get("low"), "F8696B"),
                                                                  mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                                  end_type="max", end_color=_hex(cf.get("high"), "63BE7B")))
            elif kind == "data_bar":
                ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=_hex(cf.get("color"), "5B9BD5")))
            elif kind == "cell_is":
                ws.conditional_formatting.add(rng, CellIsRule(operator=cf.get("operator", "lessThan"),
                                                              formula=[str(x) for x in (cf.get("values") or [0])],
                                                              fill=PatternFill("solid", fgColor=_hex(cf.get("fill"), "FFC7CE")),
                                                              font=Font(color=_hex(cf.get("font_color"), "9C0006"))))
            elif kind == "formula":
                ws.conditional_formatting.add(rng, FormulaRule(formula=[cf["formula"]],
                                                               fill=PatternFill("solid", fgColor=_hex(cf.get("fill"), "FFF2CC"))))
            else:
                raise OfficeError("conditional type must be color_scale, data_bar, cell_is or formula")
        for dv in sh.get("validation") or []:
            if dv.get("type", "list") == "list":
                vals = dv.get("values") or []
                v = DataValidation(type="list", formula1='"' + ",".join(map(str, vals)) + '"', allow_blank=True)
            else:
                v = DataValidation(type=dv["type"], operator=dv.get("operator", "between"),
                                   formula1=str(dv.get("min", 0)), formula2=str(dv.get("max", 0)))
            v.error, v.errorTitle = dv.get("error", "Invalid value"), "AutoMoto"
            ws.add_data_validation(v)
            v.add(dv["range"])
        for ch in sh.get("charts") or []:
            kind = ch.get("type", "bar")
            cls = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart, "scatter": ScatterChart}.get(kind)
            if not cls:
                raise OfficeError("chart type must be bar, column, line, pie or scatter")
            chart = cls()
            if kind == "bar":
                chart.type = "bar"
            elif kind == "column":
                chart.type = "col"
            chart.title = ch.get("title")
            chart.style = int(ch.get("style", 10))
            chart.height, chart.width = float(ch.get("height_cm", 7.5)), float(ch.get("width_cm", 15))
            dm = _RANGE.match(ch["data"])
            cm = _RANGE.match(ch["categories"]) if ch.get("categories") else None
            if not dm:
                raise OfficeError("chart data must be a range like B1:C6 (first row = series names)")
            col = openpyxl.utils.column_index_from_string
            data = Reference(ws, min_col=col(dm.group(1)), min_row=int(dm.group(2)),
                             max_col=col(dm.group(4) or dm.group(1)), max_row=int(dm.group(5) or dm.group(2)))
            if kind == "scatter":
                xs = Reference(ws, min_col=col(cm.group(1)), min_row=int(cm.group(2)), max_row=int(cm.group(5)))
                for c in range(data.min_col, data.max_col + 1):
                    ys = Reference(ws, min_col=c, min_row=data.min_row, max_row=data.max_row)
                    chart.series.append(Series(ys, xs, title_from_data=True))
            else:
                chart.add_data(data, titles_from_data=True)
                if cm:
                    chart.set_categories(Reference(ws, min_col=col(cm.group(1)), min_row=int(cm.group(2)),
                                                   max_col=col(cm.group(4) or cm.group(1)), max_row=int(cm.group(5) or cm.group(2))))
            palette = [_hex(c, "2E75B6") for c in (ch.get("colors") or [])] or \
                ["2E75B6", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47", "264478", "9E480E"]
            if kind in ("bar", "column", "line", "scatter"):
                # explicit colours: LibreOffice ignores openpyxl's style numbers and would draw empty outlines
                for k, ser in enumerate(chart.series):
                    col = palette[k % len(palette)]
                    if kind in ("bar", "column"):
                        ser.graphicalProperties.solidFill = col
                        ser.graphicalProperties.line.solidFill = col
                    else:
                        ser.graphicalProperties.line.solidFill = col
                        ser.graphicalProperties.line.width = 28575  # 2.25 pt
                        ser.smooth = False
                if kind in ("bar", "column"):
                    chart.gapWidth = 80
                    chart.overlap = -10 if len(chart.series) > 1 else 0
            if kind != "pie":
                chart.legend.position = "b"
                if ch.get("x_title"):
                    chart.x_axis.title = ch["x_title"]
                if ch.get("y_title"):
                    chart.y_axis.title = ch["y_title"]
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            ws.add_chart(chart, ch.get("anchor", f"{get_column_letter(c0 + ncols + 1)}{r0}"))
        if sh.get("tab_color"):
            ws.sheet_properties.tabColor = _hex(sh["tab_color"], "1F4E79")
        if sh.get("landscape", bool(sh.get("charts")) or ncols > 7):
            ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4 if str((spec.get("page") or {}).get("size", "A4")).upper() == "A4" else ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0  # fit columns to the page width, let rows flow onto more pages
        ws.print_options.horizontalCentered = True
        if header:
            ws.print_title_rows = f"{r0}:{r0}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
        report.append(f"{name}: {len(rows)} rows, {nform} formulas, {len(sh.get('charts') or [])} chart(s)")
    for nm in spec.get("names") or []:
        from openpyxl.workbook.defined_name import DefinedName

        wb.defined_names[nm["name"]] = DefinedName(nm["name"], attr_text=nm["ref"])
    wb.properties.title = str((spec.get("properties") or {}).get("title", ""))[:250]
    wb.properties.creator = str((spec.get("properties") or {}).get("author", ""))[:250]
    wb.save(str(out))
    return f"wrote {sandbox.rel(out)} - " + "; ".join(report) + ". Run office_recalc to check every formula."


# ==================================================================================== PowerPoint
def _rgb(c):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(c)


def create_pptx(spec: dict[str, Any], out: Path) -> str:
    import pptx
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    slides = spec.get("slides") if isinstance(spec, dict) else None
    if not slides or not isinstance(slides, list) or len(slides) > 300:
        raise OfficeError("spec needs slides: [{layout, title, ...}] (max 300)")
    theme = spec.get("theme") or {}
    primary = _hex(theme.get("primary"), "1F3864")
    accent = _hex(theme.get("accent"), "2E75B6")
    bg = theme.get("background")
    text_color = _hex(theme.get("text"), "262626")
    font = theme.get("font", "Calibri")
    head_font = theme.get("heading_font", font)
    prs = pptx.Presentation()
    old_w = prs.slide_width
    if spec.get("aspect", "16:9") == "16:9":
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sx = prs.slide_width / old_w
    # stretch the template's placeholders to the new width so layouts stay balanced
    # read every geometry first (layouts inherit from the master), then write explicit, scaled values
    geo = [(sh, sh.left, sh.top, sh.width, sh.height) for layout in [prs.slide_master, *prs.slide_layouts]
           for sh in layout.placeholders]
    for sh, left, top, width, height in geo:
        if None not in (left, top, width, height):
            sh.left, sh.top = Emu(int(left * sx)), top
            sh.width, sh.height = Emu(int(width * sx)), height
    W, H = prs.slide_width, prs.slide_height
    L = {"title": 0, "bullets": 1, "section": 2, "two_content": 3, "comparison": 4, "title_only": 5, "blank": 6}

    def style_text(tf, size=None, color=text_color, bold=None, fname=font):
        for para in tf.paragraphs:
            for r in para.runs:
                r.font.name = fname
                r.font.color.rgb = _rgb(color)
                if size:
                    r.font.size = Pt(size)
                if bold is not None:
                    r.font.bold = bold

    def autofit(tf):
        tf.word_wrap = True
        body = tf._txBody.find(qn("a:bodyPr"))
        for child in list(body):
            if child.tag in (qn("a:spAutoFit"), qn("a:noAutofit"), qn("a:normAutofit")):
                body.remove(child)
        from lxml import etree

        etree.SubElement(body, qn("a:normAutofit"))

    def fill_bullets(tf, items, base=24):
        tf.clear()
        n = sum(1 for _ in items)
        size = base if n <= 5 else base - 4 if n <= 7 else base - 6
        for k, item in enumerate(items):
            lvl = 0
            if isinstance(item, (list, tuple)):
                item, lvl = item[0], int(item[1])
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.level = max(0, min(lvl, 4))
            parts = re.split(r"(\*\*[^*]+\*\*)", str(item))
            for part in parts:
                if not part:
                    continue
                r = para.add_run()
                r.text = part[2:-2] if part.startswith("**") else part
                r.font.bold = part.startswith("**")
                r.font.size = Pt(size - 4 * para.level)
                r.font.name = font
                r.font.color.rgb = _rgb(text_color)
            para.space_after = Pt(8)
        autofit(tf)

    def decorate(slide, i, kind):
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(_hex(bg, "FFFFFF"))
        if kind not in ("title", "section"):
            bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(accent)
            bar.line.fill.background()
            bar.name = "accent bar"
        if spec.get("slide_numbers", True) and kind != "title":
            tb = slide.shapes.add_textbox(W - Inches(1.1), H - Inches(0.55), Inches(0.9), Inches(0.4))
            tb.name = "slide number"
            tb.text_frame.text = str(i)
            tb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
            style_text(tb.text_frame, 12, "7F7F7F")
        if spec.get("footer") and kind != "title":
            tb = slide.shapes.add_textbox(Inches(0.6), H - Inches(0.55), W - Inches(2.5), Inches(0.4))
            tb.name = "footer"
            tb.text_frame.text = str(spec["footer"])
            style_text(tb.text_frame, 11, "7F7F7F")

    def content_box(slide):
        top = Inches(1.65)
        return Inches(0.6), top, W - Inches(1.2), H - top - Inches(0.8)

    report = []
    for i, s in enumerate(slides, 1):
        if not isinstance(s, dict):
            raise OfficeError("each slide must be an object")
        kind = s.get("layout", "bullets")
        base = {"image": "title_only", "table": "title_only", "chart": "title_only", "statement": "title_only",
                "quote": "title_only"}.get(kind, kind)
        if base not in L:
            raise OfficeError(f"unknown layout '{kind}'. Use title, section, bullets, two_content, comparison, "
                              "image, table, chart, statement, quote, blank")
        slide = prs.slides.add_slide(prs.slide_layouts[L[base]])
        decorate(slide, i, kind)
        title = slide.shapes.title
        if title is not None:
            title.text_frame.text = str(s.get("title", ""))
            style_text(title.text_frame, None, primary, True, head_font)
            if kind not in ("title", "section"):
                title.top, title.height = Inches(0.35), Inches(1.1)
                title.left, title.width = Inches(0.6), W - Inches(1.2)
                title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
                style_text(title.text_frame, 34 if len(str(s.get("title", ""))) < 45 else 28, primary, True, head_font)
            autofit(title.text_frame)
        phs = {p.placeholder_format.idx: p for p in slide.placeholders}
        if kind in ("title", "section"):
            sub = phs.get(1)
            if sub is not None:
                sub.text_frame.text = str(s.get("subtitle", ""))
                style_text(sub.text_frame, 22 if kind == "title" else 20, accent, False)
                if not s.get("subtitle"):
                    sub._element.getparent().remove(sub._element)
        elif kind == "bullets":
            fill_bullets(phs[1].text_frame, s.get("bullets") or [])
        elif kind in ("two_content", "comparison"):
            left_i, right_i = (1, 2) if kind == "two_content" else (2, 4)
            if kind == "comparison":
                for idx, key in ((1, "left_title"), (3, "right_title")):
                    if idx in phs:
                        phs[idx].text_frame.text = str(s.get(key, ""))
                        style_text(phs[idx].text_frame, 20, accent, True)
            fill_bullets(phs[left_i].text_frame, s.get("left") or [], 22)
            if s.get("right_image"):
                ph = phs[right_i]
                l_, t_, w_, h_ = ph.left, ph.top, ph.width, ph.height
                ph._element.getparent().remove(ph._element)
                _add_picture(slide, _image(s["right_image"]), l_, t_, w_, h_, s.get("alt"))
            else:
                fill_bullets(phs[right_i].text_frame, s.get("right") or [], 22)
        elif kind == "image":
            l_, t_, w_, h_ = content_box(slide)
            cap_h = Inches(0.5) if s.get("caption") else 0
            _add_picture(slide, _image(s.get("image", "")), l_, t_, w_, h_ - cap_h, s.get("alt") or s.get("caption"))
            if s.get("caption"):
                tb = slide.shapes.add_textbox(l_, t_ + h_ - cap_h, w_, cap_h)
                tb.text_frame.text = str(s["caption"])
                tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                style_text(tb.text_frame, 14, "595959")
        elif kind == "table":
            rows = s.get("rows") or []
            if not rows:
                raise OfficeError("table slide needs rows")
            l_, t_, w_, h_ = content_box(slide)
            nr, nc = len(rows), max(len(r) for r in rows)
            rh = min(Inches(0.5), int(h_ / nr))
            gf = slide.shapes.add_table(nr, nc, l_, t_, w_, rh * nr)
            tbl = gf.table
            fs = 16 if nr <= 6 else 13 if nr <= 10 else 11
            for r_i, r in enumerate(rows):
                for c_i in range(nc):
                    cell = tbl.cell(r_i, c_i)
                    cell.text = str(r[c_i]) if c_i < len(r) and r[c_i] is not None else ""
                    style_text(cell.text_frame, fs, "FFFFFF" if r_i == 0 else text_color, r_i == 0)
                    if r_i == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = _rgb(primary)
        elif kind == "chart":
            ch = s.get("chart") or {}
            cd = CategoryChartData()
            cd.categories = list(ch.get("categories") or [])
            for ser in ch.get("series") or []:
                cd.add_series(str(ser.get("name", "")), [float(v) for v in ser.get("values", [])])
            ctype = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                     "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
                     "stacked": XL_CHART_TYPE.COLUMN_STACKED}.get(ch.get("type", "column"))
            if ctype is None:
                raise OfficeError("chart type must be bar, column, line, pie or stacked")
            l_, t_, w_, h_ = content_box(slide)
            gf = slide.shapes.add_chart(ctype, l_, t_, w_, h_, cd)
            chart = gf.chart
            chart.has_legend = len(ch.get("series") or []) > 1 or ch.get("type") == "pie"
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                chart.legend.font.size = Pt(14)
            chart.font.size = Pt(14)
            chart.font.name = font
            if ch.get("data_labels", True):
                plot = chart.plots[0]
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(12)
                if ch.get("number_format"):
                    plot.data_labels.number_format = ch["number_format"]
                    plot.data_labels.number_format_is_linked = False
            if ch.get("type") != "pie":
                for k, ser in enumerate(chart.plots[0].series):
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = _rgb([accent, primary, "A5A5A5", "FFC000", "70AD47"][k % 5])
        elif kind in ("statement", "quote"):
            l_, t_, w_, h_ = content_box(slide)
            tb = slide.shapes.add_textbox(l_ + Inches(0.5), t_ + Inches(0.4), w_ - Inches(1), h_ - Inches(0.8))
            tb.text_frame.word_wrap = True
            txt = str(s.get("text", ""))
            tb.text_frame.text = (f"“{txt}”" if kind == "quote" else txt)
            tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            style_text(tb.text_frame, 36 if len(txt) < 80 else 28, primary, kind == "statement")
            if s.get("attribution"):
                p = tb.text_frame.add_paragraph()
                p.text = "— " + str(s["attribution"])
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size, r.font.color.rgb, r.font.name = Pt(20), _rgb("595959"), font
            autofit(tb.text_frame)
        # unused empty placeholders confuse readers and screen readers - remove them
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
        report.append(kind)
    prs.core_properties.title = str((spec.get("properties") or {}).get("title") or slides[0].get("title", ""))[:250]
    prs.core_properties.author = str((spec.get("properties") or {}).get("author", ""))[:250]
    prs.save(str(out))
    with_notes = sum(1 for s in slides if s.get("notes"))
    return f"wrote {sandbox.rel(out)}: {len(slides)} slides ({', '.join(report)}), speaker notes on {with_notes}"


def _add_picture(slide, img: Path, left, top, width, height, alt=None):
    from PIL import Image

    with Image.open(img) as im:
        iw, ih = im.size
    scale = min(width / iw, height / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(str(img), left + (width - w) // 2, top + (height - h) // 2, w, h)
    pic._element.nvPicPr.cNvPr.set("descr", str(alt or img.stem)[:500])
    return pic


# ======================================================================================== edit
def edit(path: Path, replace: list[dict[str, str]] | None = None, cells: dict[str, Any] | None = None,
         append: list[Any] | None = None, notes: dict[str, str] | None = None) -> str:
    ext = path.suffix.lower()
    inspect_file(path)
    changes = 0
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        paras = list(d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    paras.extend(c.paragraphs)
        for sec in d.sections:
            paras.extend(sec.header.paragraphs)
            paras.extend(sec.footer.paragraphs)
        for r in replace or []:
            old, new = str(r["old"]), str(r["new"])
            for p in paras:
                if old not in p.text:
                    continue
                done = False
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        changes += 1
                        done = True
                if not done and p.runs:  # text split across runs: merge into the first run
                    full = p.text.replace(old, new)
                    p.runs[0].text = full
                    for run in p.runs[1:]:
                        run.text = ""
                    changes += 1
        if append:
            tmp = Path(tempfile.mkdtemp(prefix="automoto-app-")) / "a.docx"
            build_docx({"blocks": append, "page_numbers": False}, tmp)
            src = docx.Document(str(tmp))
            body = d.element.body
            sect = body[-1] if body[-1].tag.endswith("sectPr") else None
            for el in list(src.element.body):
                if el.tag.endswith("sectPr"):
                    continue
                if sect is not None:
                    sect.addprevious(el)
                else:
                    body.append(el)
                changes += 1
        d.save(str(path))
    elif ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(path), keep_links=False)
        for ref, v in (cells or {}).items():
            sheet, _, addr = ref.rpartition("!")
            ws = wb[sheet.strip("'")] if sheet else wb.active
            ws[addr] = v
            changes += 1
        for r in replace or []:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and r["old"] in c.value:
                            c.value = c.value.replace(r["old"], r["new"])
                            changes += 1
        wb.save(str(path))
    elif ext == ".pptx":
        import pptx

        prs = pptx.Presentation(str(path))
        for r in replace or []:
            for s in prs.slides:
                for sh in s.shapes:
                    if sh.has_text_frame:
                        for p in sh.text_frame.paragraphs:
                            for run in p.runs:
                                if r["old"] in run.text:
                                    run.text = run.text.replace(r["old"], r["new"])
                                    changes += 1
        for k, txt in (notes or {}).items():
            idx = int(k) - 1
            if not 0 <= idx < len(prs.slides):
                raise OfficeError(f"no slide {k}")
            prs.slides[idx].notes_slide.notes_text_frame.text = str(txt)
            changes += 1
        prs.save(str(path))
    else:
        raise OfficeError("office_edit works on .docx, .xlsx and .pptx (convert other formats first)")
    return f"edited {sandbox.rel(path)}: {changes} change(s)"
