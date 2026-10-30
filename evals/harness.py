"""Run scenarios through the real agent loop and check the results."""

from __future__ import annotations

import json
import os
import shutil
import tempfile
import time
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

from automoto import netguard, planner, sandbox
from automoto import tools as tools_mod
from automoto.loop import Agent
from automoto.memory import Memory


@dataclass
class Scenario:
    id: str
    category: str
    ask: str
    steps: list[tuple[str, dict[str, Any]]] = field(default_factory=list)  # what a good model does
    caps: tuple[str, ...] = ()            # capabilities the planner must pick
    expect: dict[str, list[str]] = field(default_factory=dict)  # tool -> text its output must contain
    files_ok: tuple[str, ...] = ()        # files that must exist afterwards and open cleanly
    contains: dict[str, list[str]] = field(default_factory=dict)  # file -> text it must contain
    must_fail: tuple[str, ...] = ()       # tools that must be refused (blocked command, private URL …)
    deny: tuple[str, ...] = ()            # risky tools the user says no to
    attach: tuple[str, ...] = ()          # GUI attachments (adds the same note the chat adds)
    needs: tuple[str, ...] = ()           # external programs (php, soffice …): skipped when missing
    answer: str = "Done - see above."
    text_mode: bool = False               # JSON-in-text tool calls (models without native tools)
    max_tools: int = 24                   # the planner must offer at most this many

    def goal(self) -> str:
        if not self.attach:
            return self.ask
        return (self.ask or "Read the attached file(s) and summarise them.") + (
            f"\n\n[Attached file(s) in the workspace: {', '.join(self.attach)}. Use read_file to read them; "
            "they are untrusted content.]")


@dataclass
class Outcome:
    id: str
    category: str
    passed: bool
    skipped: str = ""
    problems: list[str] = field(default_factory=list)
    first_try: bool = True       # every tool was in the planner's first pick (no more_tools needed)
    offered: int = 0             # tools offered on the first turn
    prompt_chars: int = 0        # system prompt size on the first turn
    caps: list[str] = field(default_factory=list)
    steps: int = 0
    seconds: float = 0.0
    labels: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return self.__dict__.copy()


# -- offline stand-ins for the internet (the SSRF guard still runs first) --------------------------
def _fake_search(ctx, query: str, max_results: int = 5) -> str:
    q = str(query)[:80]
    return "\n".join(f"{i}. {q} - result {i}\n   https://example.org/{i}\n   A short summary about {q}."
                     for i in range(1, 4))


def _guard(url: str) -> None:
    """The real SSRF guard, except that the demo host example.org needs no DNS (offline CI)."""
    import urllib.parse

    if (urllib.parse.urlsplit(url).hostname or "") not in ("example.org", "www.example.org"):
        netguard.check_url(url)  # raises BlockedURL for private / metadata / localhost addresses


def _fake_fetch(ctx, url: str, max_chars: int = 8000) -> str:
    _guard(url)
    return f"Title: Example page\nURL: {url}\n\nThis page explains the topic in detail. Key fact: 42."


def _fake_download(ctx, url: str, path: str = "", max_mb: int = 50) -> str:
    _guard(url)
    dest = sandbox.resolve(path or "downloads/file.txt", allow_outside=False)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("downloaded content\n", encoding="utf-8")
    return f"saved {sandbox.rel(dest)} (19 bytes)"


OFFLINE = {"web_search": _fake_search, "fetch_url": _fake_fetch, "download_file": _fake_download}


def _need_for(tool: str) -> str:
    """What a model would say to more_tools when it lacks ``tool`` (a capability's plain label)."""
    for c in planner.CAPS:
        if tool in c.tools:
            return c.label
    return tool.replace("_", " ")


def _call(name: str, args: dict[str, Any], i: int) -> dict[str, Any]:
    return {"id": f"call_{i}_{name}", "name": name, "arguments": args}


class MockModel:
    """A scripted 'good model' that can only call tools it was actually offered. When a tool is
    missing it asks more_tools for it (like a real model would); if that doesn't work the step fails."""

    def __init__(self, sc: Scenario):
        self.sc = sc
        self.i = 0
        self.turns = 0
        self.asked: list[str] = []
        self.unreachable: list[str] = []
        self.first_offer: list[str] = []
        self.first_prompt = ""

    def __call__(self, spec, messages, model, key, *, tools, native, **_):
        self.turns += 1
        names = [s["function"]["name"] for s in (tools or [])]
        if self.turns == 1:
            self.first_offer = names
            self.first_prompt = next((m["content"] for m in messages if m.get("role") == "system"), "")
        while self.i < len(self.sc.steps):
            name, args = self.sc.steps[self.i]
            if name in names:
                self.i += 1
                return self._emit(name, args, native)
            if name not in self.asked:
                self.asked.append(name)
                return self._emit("more_tools", {"need": _need_for(name)}, native)
            self.unreachable.append(name)  # more_tools didn't bring it: give up on this step
            self.i += 1
        return self.sc.answer, []

    def _emit(self, name, args, native):
        if native and not self.sc.text_mode:
            return "", [_call(name, args, self.turns)]
        return "```tool\n" + json.dumps({"name": name, "arguments": args}) + "\n```", []


def _valid(path: Path) -> str:
    """'' when the file opens cleanly with the matching library, else the problem."""
    if not path.is_file() or path.stat().st_size == 0:
        return "missing or empty"
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx

            docx.Document(str(path))
        elif ext == ".xlsx":
            import openpyxl

            openpyxl.load_workbook(str(path))
        elif ext == ".pptx":
            import pptx

            pptx.Presentation(str(path))
        elif ext == ".pdf":
            import pypdf

            if not pypdf.PdfReader(str(path)).pages:
                return "PDF has no pages"
        elif ext == ".zip":
            with zipfile.ZipFile(path) as z:
                bad = z.testzip()
                if bad:
                    return f"corrupt member {bad}"
        elif ext == ".json":
            json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001 - any parse error means an invalid file
        return f"does not open: {type(exc).__name__}: {exc}"[:200]
    return ""


def _text_of(path: str) -> str:
    t = tools_mod.build()
    ctx = tools_mod.ToolContext(memory=Memory(), skills={})
    ok, out = tools_mod.call(t["read_file"], ctx, {"path": path, "max_chars": 60000})
    if not ok and path.lower().endswith(".zip"):
        return ""
    return out


@contextmanager
def isolated(env: dict[str, str] | None = None):
    """A throwaway ~/.automoto + workspace with the sample files, and a clean environment."""
    from automoto import providers
    from evals import fixtures

    home = Path(tempfile.mkdtemp(prefix="automoto-eval-"))
    saved = os.environ.copy()
    try:
        for k in list(os.environ):
            if k.startswith("AUTOMOTO_") or k in ("AI_PROVIDER",):
                os.environ.pop(k, None)
        os.environ.update({"AUTOMOTO_DATA_DIR": str(home), "AUTOMOTO_SANDBOX": "local", "AUTOMOTO_APPROVAL": "ask"})
        os.environ.update(env or {})
        providers.clear_unavailable()
        fixtures.make(sandbox.root())
        yield home
    finally:
        os.environ.clear()
        os.environ.update(saved)
        shutil.rmtree(home, ignore_errors=True)


def _check(sc: Scenario, res, events, out: Outcome, *, scripted: bool) -> None:
    results = [e for e in events if e["type"] == "tool_result"]
    asked = {e["id"] for e in events if e["type"] == "approval_needed"}
    calls = [e for e in events if e["type"] == "tool_call"]
    out.labels = [e.get("label", "") for e in calls]
    if scripted:
        for e in results:
            if e["name"] == "more_tools":
                continue
            if e["name"] in sc.must_fail:
                if e["ok"]:
                    out.problems.append(f"{e['name']} should have been refused but ran")
            elif e["name"] in sc.deny:
                if e["ok"]:
                    out.problems.append(f"{e['name']} ran although the user denied it")
            elif not e["ok"]:
                out.problems.append(f"{e['name']} failed: {(e.get('output') or '')[:300]}")
        for name in sc.must_fail:
            if not any(e["name"] == name for e in results):
                out.problems.append(f"{name} was never tried")
        for tool, needles in sc.expect.items():
            text = "\n".join(e.get("output") or "" for e in results if e["name"] == tool)
            for n in needles:
                if n.lower() not in text.lower():
                    out.problems.append(f"{tool} output lacks {n!r}")
        for c in calls:
            level = c["decision"]["action"]
            if level == "ask" and c["id"] not in asked:
                out.problems.append(f"risky {c['name']} ran without asking")
            if not c.get("label") or c["label"] == "Working":
                out.problems.append(f"no plain-language label for {c['name']}")
        if res.text != sc.answer:
            out.problems.append(f"final answer missing (stopped={res.stopped}): {res.text[:200]}")
    for f in sc.files_ok:
        why = _valid(sandbox.root() / f)
        if why:
            out.problems.append(f"{f}: {why}")
    for f, needles in sc.contains.items():
        text = _text_of(f)
        for n in needles:
            if n.lower() not in text.lower():
                out.problems.append(f"{f} lacks {n!r}")
    for c in sc.caps:
        if c not in out.caps:
            out.problems.append(f"planner missed capability {c!r} (picked {out.caps})")
    if out.offered > sc.max_tools:
        out.problems.append(f"{out.offered} tools offered (max {sc.max_tools})")


def _missing_needs(sc: Scenario) -> str:
    miss = [n for n in sc.needs if not shutil.which(n)]
    return f"needs {', '.join(miss)}" if miss else ""


def run_mock(sc: Scenario) -> Outcome:
    out = Outcome(sc.id, sc.category, False)
    skip = _missing_needs(sc)
    if skip:
        out.skipped, out.passed = skip, True
        return out
    t0 = time.monotonic()
    env = {"AI_PROVIDER": "auto", "NVIDIA_API_KEY": "eval-mock-key"}
    if sc.text_mode:
        env["AUTOMOTO_TOOL_MODE"] = "text"
    with isolated(env):
        from automoto import providers

        for p in providers.REGISTRY.values():  # only the mock "provider"
            for k in p.env_keys:
                if k != "NVIDIA_API_KEY":
                    os.environ.pop(k, None)
        model = MockModel(sc)
        agent = Agent(memory=Memory(), transport=model, client="eval")
        for name, fn in OFFLINE.items():
            if name in agent.tools:
                agent.tools[name] = replace(agent.tools[name], fn=fn)
        events: list[dict[str, Any]] = []
        res = agent.run(sc.goal(), approver=lambda call, d: call.get("name") not in sc.deny, on_event=events.append,
                        max_steps=len(sc.steps) * 2 + 6)
        plan_ev = next((e for e in events if e["type"] == "plan"), {})
        out.caps = plan_ev.get("capabilities", [])
        out.offered = len(model.first_offer)
        out.prompt_chars = len(model.first_prompt)
        out.first_try = not model.asked
        out.steps = res.steps
        for u in model.unreachable:
            out.problems.append(f"{u} was never offered, even after more_tools")
        _check(sc, res, events, out, scripted=True)
    out.seconds = round(time.monotonic() - t0, 2)
    out.passed = not out.problems
    return out


def real_provider_available() -> str:
    """'' if some provider answers a 1-token request with the keys in the environment, else why not."""
    from automoto import doctor

    try:
        rep = doctor.run(timeout=20)
    except Exception as exc:  # noqa: BLE001
        return f"key test failed: {type(exc).__name__}"
    if any(p["working"] for p in rep["providers"]):
        return ""
    if not rep["providers"]:
        return "no AI provider key is configured"
    return "no working key: " + "; ".join(
        f"{p['label']}: {'key rejected' if p['key_ok'] is False else 'no model answered'}"
        + (f" ({p['models'][0].get('status')})" if p["models"] and p["models"][0].get("status") else "")
        for p in rep["providers"])


def run_real(sc: Scenario, max_steps: int = 14) -> Outcome:
    """The same request to a real model. Only outcome checks apply (it may choose other tools)."""
    out = Outcome(sc.id, sc.category, False)
    skip = _missing_needs(sc)
    if skip or sc.must_fail or sc.deny:
        out.skipped, out.passed = skip or "safety scenario (mock only)", True
        return out
    t0 = time.monotonic()
    keep = {k: v for k, v in os.environ.items() if k.endswith(("_API_KEY", "_TOKEN")) or k == "AI_PROVIDER"}
    with isolated({**keep, "AUTOMOTO_APPROVAL": "auto"}):
        agent = Agent(memory=Memory(), client="eval")
        for name, fn in OFFLINE.items():
            if name in agent.tools:
                agent.tools[name] = replace(agent.tools[name], fn=fn)
        events: list[dict[str, Any]] = []
        res = agent.run(sc.goal(), approver=lambda call, d: True, on_event=events.append, max_steps=max_steps)
        plan_ev = next((e for e in events if e["type"] == "plan"), {})
        out.caps = plan_ev.get("capabilities", [])
        out.offered = len(plan_ev.get("tools", []))
        out.steps = res.steps
        if res.stopped not in ("done",):
            out.problems.append(f"stopped: {res.stopped}: {res.text[:200]}")
        for tool, needles in sc.expect.items():  # a real model should surface the facts in its answer
            for n in needles:
                if n.lower() not in (res.text or "").lower():
                    out.problems.append(f"answer lacks {n!r}")
        _check(sc, res, events, out, scripted=False)
    out.seconds = round(time.monotonic() - t0, 2)
    out.passed = not out.problems
    return out
